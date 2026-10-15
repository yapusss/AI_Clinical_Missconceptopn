"""Create a question bank (QuestionSet + Questions + Versions) from a TOPIK .docx.

The web import flow (`POST /api/question-imports`) only *parses and previews* — it never
persists an `ImportJob`, and `QuestionImportCommitView` requires one in `READY_TO_IMPORT`.
Nothing in the codebase creates one, so that commit endpoint can never succeed. This command
is the working path, and it is reusable for TOPIK 2 and 3.

    manage.py import_question_bank --docx "<path>" --subject Physics --publish
    manage.py import_question_bank --docx "<path>" --dry-run
    manage.py import_question_bank --docx "<path>" --xlsx-out D:/out/bank.xlsx

Mapping (from `api.question_bank`):
    Soal            -> QuestionVersion.prompt
    Jawaban         -> QuestionVersion.short_answer
    Alasan jawaban  -> QuestionVersion.model_answer   (+ the 2-row reference_answers mirror)

`order_index` is renumbered 1..N because the app stores it as an integer and the source
numbers include `5.1` / `18.2`, which would both collide after `int(float(...))`. The source
number is preserved verbatim in `Question.external_key`.
"""
import os
import uuid

from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from api.models import (Question, QuestionSet, QuestionVersion, Subject, Topic, User)
from api.question_bank import parse_instrument_docx, row_problems
from api.views import _sync_reference_answers


class Command(BaseCommand):
    help = 'Create a question bank from a TOPIK instrument .docx'

    def add_arguments(self, parser):
        parser.add_argument('--docx', required=True, help='path to the .docx instrument')
        parser.add_argument('--subject', default='Physics')
        parser.add_argument('--topic', default='', help='defaults to the topic declared in the docx')
        parser.add_argument('--code', default='', help='defaults to FIS-TOPIK-<n> or a random code')
        parser.add_argument('--title', default='')
        parser.add_argument('--description', default='')
        parser.add_argument('--created-by', default='lecturer@acm.local')
        parser.add_argument('--publish', action='store_true',
                            help='publish the versions (the web import leaves them unpublished)')
        parser.add_argument('--dry-run', action='store_true')
        parser.add_argument('--xlsx-only', action='store_true',
                            help='only write the import XLSX; do not touch the DB')
        parser.add_argument('--replace', action='store_true',
                            help='delete an existing question set with the same code first')
        parser.add_argument('--xlsx-out', default='', help='also write the import-format XLSX here')

    def handle(self, *args, **o):
        path = o['docx']
        if not os.path.exists(path):
            raise CommandError(f'file not found: {path}')

        declared_topic, rows = parse_instrument_docx(path)
        topic_name = (o['topic'] or declared_topic).strip()
        if not topic_name:
            raise CommandError('no topic: pass --topic or use a docx that declares one')

        code = o['code'] or self._derive_code(path, declared_topic)
        title = o['title'] or declared_topic or topic_name
        description = o['description'] or f'Diimpor dari {os.path.basename(path)}'

        # ---- report the parse before touching anything ------------------------
        self.stdout.write(f'source       : {os.path.basename(path)}')
        self.stdout.write(f'declared topic: {declared_topic!r}')
        self.stdout.write(f'topic        : {topic_name!r}')
        self.stdout.write(f'code/title   : {code} / {title!r}')
        self.stdout.write(f'rows parsed  : {len(rows)}')
        n_img = sum(1 for r in rows if r['images'])
        self.stdout.write(f'rows w/ figure: {n_img}  '
                          f'({sum(len(r["images"]) for r in rows)} images)')

        incomplete = [(r['no_soal'], row_problems(r)) for r in rows if row_problems(r)]
        if incomplete:
            self.stdout.write(self.style.WARNING(
                f'incomplete rows ({len(incomplete)}):'))
            for no_soal, probs in incomplete:
                self.stdout.write(self.style.WARNING(f'   Soal {no_soal}: {", ".join(probs)}'))
        else:
            self.stdout.write('incomplete rows: none')

        if o['xlsx_only']:
            subject = Subject.objects.filter(name__iexact=o['subject'], is_active=True).first()
            if not subject:
                raise CommandError(f'subject not found: {o["subject"]}')
            written = self._write_xlsx(o['xlsx_out'] or f'{code}_bank.xlsx', subject,
                                       topic_name, code, title, description, rows)
            self.stdout.write(self.style.SUCCESS(f'wrote import XLSX: {written}'))
            return

        if o['dry_run']:
            self.stdout.write(self.style.WARNING('--dry-run: nothing written'))
            return

        # ---- guards ----------------------------------------------------------
        subject = Subject.objects.filter(name__iexact=o['subject'], is_active=True).first()
        if not subject:
            raise CommandError(f'subject not found: {o["subject"]}')
        creator = User.objects.filter(email=o['created_by']).first()
        if not creator:
            raise CommandError(f'user not found: {o["created_by"]}')
        existing = QuestionSet.objects.filter(code__iexact=code).first()
        if existing and not o['replace']:
            raise CommandError(f'question set code already exists: {code} '
                               f'(pass --replace to delete and recreate it)')
        if existing:
            n_q = Question.objects.filter(question_set=existing).count()
            n_v = QuestionVersion.objects.filter(question__question_set=existing).count()
            self.stdout.write(self.style.WARNING(
                f'--replace: deleting {code} ({n_q} questions, {n_v} versions)'))
            existing.delete()

        with transaction.atomic():
            topic = Topic.objects.filter(subject=subject, name__iexact=topic_name).first()
            topic_created = False
            if not topic:
                topic = Topic.objects.create(id=uuid.uuid4(), subject=subject,
                                             name=topic_name,
                                             description='Dibuat otomatis dari impor bank soal.')
                topic_created = True

            q_set = QuestionSet.objects.create(
                id=uuid.uuid4(), subject=subject, topic=topic, created_by=creator,
                code=code, title=title, description=description, is_active=True,
            )

            created, with_img = 0, 0
            for order_index, row in enumerate(rows, start=1):
                prompt = row['soal'].strip()
                if row['images']:
                    # Mirror `QuestionImportCreateView`: append the figure to the prompt as
                    # HTML. Prompts without a figure stay plain text so the raw prompt the
                    # LLM reads (load_question_version_context) carries no markup.
                    parts = []
                    for img in row['images']:
                        ext = os.path.splitext(img['name'])[1] or '.png'
                        saved = default_storage.save(
                            f'uploads/question-bank/{uuid.uuid4().hex}{ext}',
                            ContentFile(img['data']))
                        from django.conf import settings
                        url = f'{settings.MEDIA_URL.rstrip("/")}/{saved.lstrip("/")}'
                        parts.append(f'<img src="{url}" style="max-width:100%;height:auto;" />')
                    prompt = f'<p>{prompt}</p>' + ''.join(f'<p>{p}</p>' for p in parts)
                    with_img += 1

                question = Question.objects.create(
                    id=uuid.uuid4(), question_set=q_set,
                    external_key=(row['no_soal'].strip() or f'Q-{order_index}'),
                    order_index=order_index,
                )
                version = QuestionVersion.objects.create(
                    id=uuid.uuid4(), question=question, version_number=1,
                    prompt=prompt, short_answer=row['jawaban'].strip(),
                    model_answer=row['alasan'].strip(),
                    is_published=bool(o['publish']), created_by=creator,
                )
                _sync_reference_answers(version, row['jawaban'].strip(), row['alasan'].strip())
                created += 1

        self.stdout.write(self.style.SUCCESS(
            f'created question set {code} ({q_set.id})'))
        self.stdout.write(f'  topic       : {topic.name}'
                          f'{" (created)" if topic_created else " (existing)"}')
        self.stdout.write(f'  questions   : {created}')
        self.stdout.write(f'  with figure : {with_img}')
        self.stdout.write(f'  published   : {bool(o["publish"])}')

        if o['xlsx_out']:
            written = self._write_xlsx(o['xlsx_out'], subject, topic_name, code, title,
                                       description, rows)
            self.stdout.write(self.style.SUCCESS(f'wrote import XLSX: {written}'))

    # ------------------------------------------------------------------ helpers
    @staticmethod
    def _derive_code(path, declared_topic):
        import re
        m = re.search(r'\(TOPIK\s*(\d+)\)', os.path.basename(path), re.IGNORECASE)
        return f'FIS-TOPIK-{m.group(1)}' if m else f'FIS-IMP-{uuid.uuid4().hex[:6].upper()}'

    def _write_xlsx(self, out_path, subject, topic_name, code, title, description, rows):
        """Emit the file the app's own importer expects, so it previews identically."""
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = subject.name[:31]

        ws['A1'] = (f'⚠️ SHEET MATA KULIAH: {subject.name.upper()}\n'
                    'Jangan mengubah nama sheet. Baris header di bawah tidak boleh diubah.')
        ws['A1'].font = Font(bold=True, color='92400E')
        ws['A1'].fill = PatternFill('solid', start_color='FEF3C7')

        headers = ['NOMOR_SOAL', 'KODE_PAKET_UNIK', 'JUDUL_UJIAN', 'TOPIK',
                   'DESKRIPSI_INSTRUKSI', 'PERTANYAAN', 'GAMBAR', 'JAWABAN_SINGKAT',
                   'ALASAN_JAWABAN', 'NO_SOAL_ASLI', 'LEVEL_SOAL', 'TUJUAN_PEMBELAJARAN']
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=5, column=c, value=h)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', start_color='00288E')

        skipped, extra_figures = [], []
        for i, row in enumerate(rows, start=1):
            r = 5 + i
            ws.cell(row=r, column=1, value=i)                    # NOMOR_SOAL -> order_index
            ws.cell(row=r, column=2, value=code)
            ws.cell(row=r, column=3, value=title)
            ws.cell(row=r, column=4, value=topic_name)
            ws.cell(row=r, column=5, value=description)
            ws.cell(row=r, column=6, value=row['soal'])
            ws.cell(row=r, column=7, value=None)                 # GAMBAR: images are embedded
            ws.cell(row=r, column=8, value=row['jawaban'])
            ws.cell(row=r, column=9, value=row['alasan'])
            ws.cell(row=r, column=10, value=row['no_soal'])
            ws.cell(row=r, column=11, value=row['level'])
            ws.cell(row=r, column=12, value=row['tujuan'])
            for c in (6, 8, 9, 11, 12):
                ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical='top')
            # anchor each figure to its own row: the parser reads the image's row anchor
            for j, img in enumerate(row['images']):
                try:
                    xl = XLImage(BytesIO(img['data']))
                except Exception as exc:      # noqa: BLE001
                    # Never swallow this silently: openpyxl needs Pillow to embed an
                    # image, and a bare `continue` once hid 9 figures from the output.
                    skipped.append((i, img['name'], f'{type(exc).__name__}: {exc}'))
                    continue
                if j == 0:
                    # Anchor to THIS row's G cell. The app's importer maps one image per
                    # row (`images_by_row[row]`), and anchoring at G{r+j} made a row's 2nd
                    # and 3rd figures spill onto the FOLLOWING rows' anchors.
                    xl.width, xl.height = 160, 160
                    ws.add_image(xl, f'G{r}')
                else:
                    extra_figures.append((i, row['no_soal'], img['name']))

        for col, width in (('A', 11), ('B', 16), ('C', 30), ('D', 30), ('E', 30),
                           ('F', 70), ('G', 12), ('H', 50), ('I', 60),
                           ('J', 12), ('K', 12), ('L', 50)):
            ws.column_dimensions[col].width = width

        if skipped:
            self.stdout.write(self.style.ERROR(
                f'   !! {len(skipped)} figure(s) NOT embedded: {skipped[:3]}'))
        if extra_figures:
            # The app's import format carries ONE image per question, so a row with
            # several figures can only place its first here. All of them live in the DB
            # bank (as HTML in the prompt).
            self.stdout.write(self.style.WARNING(
                f'   note: {len(extra_figures)} further figure(s) not placed — the import '
                f'format holds one image per row: {extra_figures}'))
        os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
        wb.save(out_path)
        return out_path

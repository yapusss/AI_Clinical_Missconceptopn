import csv
import io
import json
import uuid
from decimal import Decimal
from django.utils.text import slugify
from django.db import connection, transaction
from django.db.models import Avg, Count, Q, Sum
from django.utils import timezone
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.views import APIView
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .authentication import TokenAuthentication
from .models import (
    AuthToken,
    ConceptIndicator,
    HelpArticle,
    LlmAnalysis,
    Misconception,
    Question,
    QuestionSet,
    QuestionVersion,
    ImportJob,
    ImportRow,
    ReferenceAnswer,
    Subject,
    Submission,
    Topic,
    User,
    UserRole,
    UserSubjectRole,
    Validation,
    WebsiteSection,
)
from .serializers import (
    AdminManagedUserSerializer,
    LoginSerializer,
    QuestionSetCreateSerializer,
    QuestionSetUpdateSerializer,
    PackageSubmissionCreateSerializer,
    RegisterSerializer,
    SubmissionCreateSerializer,
    UserSerializer,
)


class RegisterView(APIView):
    authentication_classes = []

    def post(self, request):
        serializer = RegisterSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        user = serializer.save()
        token = AuthToken.generate(user)
        return Response({
            'token': token.key,
            'user': UserSerializer(user).data,
        }, status=status.HTTP_201_CREATED)


class LoginView(APIView):
    authentication_classes = []

    def post(self, request):
        serializer = LoginSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        user = serializer.validated_data['user']
        token = AuthToken.generate(user)
        return Response({
            'token': token.key,
            'user': UserSerializer(user).data,
        })


class MeView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request):
        data = UserSerializer(request.user).data
        data['roles'] = self._roles(request.user)
        return Response(data)

    def post(self, request):
        AuthToken.objects.filter(user=request.user).delete()
        return Response(status=status.HTTP_204_NO_CONTENT)

    @staticmethod
    def _roles(user):
        qs = UserRole.objects.filter(user=user)
        return [
            {
                'role': r.role,
                'subject_slug': None,
                'subject_name': None,
            }
            for r in qs
        ]


class DashboardView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        role_rows = UserSubjectRole.objects.filter(
            user=user, role=UserSubjectRole.Role.LECTURER,
        ).select_related('subject')
        roles = MeView._roles(user)
        role_set = {r['role'] for r in roles}

        summary = {
            'my_subjects': [
                {
                    'id': str(r.subject.id),
                    'slug': r.subject.slug,
                    'name': r.subject.name,
                }
                for r in role_rows
            ],
        }

        my_subjects_ids = [r.subject_id for r in role_rows]

        if user.is_superuser:
            question_sets = list(QuestionSet.objects.select_related('subject').order_by('-updated_at'))
            question_counts = {
                row['question_set_id']: row['count']
                for row in Question.objects.values('question_set_id').annotate(count=Count('id'))
            }
            published_counts = {
                row['question__question_set_id']: row['count']
                for row in QuestionVersion.objects.filter(is_published=True).values('question__question_set_id').annotate(count=Count('question_id', distinct=True))
            }
            package_status = {'DRAFT': 0, 'ACTIVE': 0, 'INACTIVE': 0}
            attention_packages = []
            for question_set in question_sets:
                question_count = question_counts.get(question_set.id, 0)
                is_published = question_count > 0 and published_counts.get(question_set.id, 0) >= question_count
                if not is_published:
                    package_status['DRAFT'] += 1
                    attention_packages.append({
                        'id': str(question_set.id),
                        'code': question_set.code,
                        'title': question_set.title,
                        'subject_name': question_set.subject.name,
                        'reason': 'Belum lengkap atau belum diterbitkan',
                        'state': 'DRAFT',
                    })
                elif question_set.is_active:
                    package_status['ACTIVE'] += 1
                else:
                    package_status['INACTIVE'] += 1

            recent_submissions = list(Submission.objects.select_related('student').order_by('-submitted_at')[:5])
            validation_status = {
                'PENDING': Submission.objects.filter(status='PENDING_VALIDATION').count(),
                'VALIDATED': Submission.objects.filter(status='VALIDATED').count(),
                'FAILED': Submission.objects.filter(status='ANALYSIS_FAILED').count(),
            }
            analyzed_count = Submission.objects.filter(status__in=['PENDING_VALIDATION', 'VALIDATED']).count()
            failed_count = Submission.objects.filter(status='ANALYSIS_FAILED').count()
            average_execution = LlmAnalysis.objects.filter(is_current=True).aggregate(value=Avg('execution_time_ms'))['value']
            activities = []
            activities.extend({
                'label': f'Akun {account.full_name} dibuat',
                'timestamp': account.created_at.isoformat(),
                'type': 'ACCOUNT',
            } for account in User.objects.order_by('-created_at')[:4])
            activities.extend({
                'label': f'Paket {question_set.code} diperbarui',
                'timestamp': question_set.updated_at.isoformat(),
                'type': 'PACKAGE',
            } for question_set in question_sets[:4])
            activities.extend({
                'label': f'Jawaban dikumpulkan oleh {submission.student.full_name}',
                'timestamp': submission.submitted_at.isoformat(),
                'type': 'SUBMISSION',
            } for submission in recent_submissions[:4])
            activities.sort(key=lambda activity: activity['timestamp'], reverse=True)

            alerts = []
            if package_status['DRAFT']:
                alerts.append({'level': 'warning', 'message': f"{package_status['DRAFT']} paket belum siap diterbitkan."})
            if validation_status['PENDING']:
                alerts.append({'level': 'info', 'message': f"{validation_status['PENDING']} jawaban menunggu validasi dosen."})
            if failed_count:
                alerts.append({'level': 'error', 'message': f'{failed_count} analisis AI gagal dan perlu diproses ulang.'})

            summary.update({
                'total_users': User.objects.count(),
                'total_subjects': Subject.objects.count(),
                'total_question_sets': QuestionSet.objects.count(),
                'total_submissions': Submission.objects.count(),
                'total_misconceptions': Misconception.objects.count(),
                'total_analyses': LlmAnalysis.objects.count(),
                'total_validations': Validation.objects.count(),
                'admin_dashboard': {
                    'system': {
                        'active_students': UserRole.objects.filter(role=UserRole.Role.STUDENT, user__is_active=True).count(),
                        'active_lecturers': UserRole.objects.filter(role=UserRole.Role.LECTURER, user__is_active=True).count(),
                        'active_subjects': Subject.objects.filter(is_active=True).count(),
                        'active_question_sets': package_status['ACTIVE'],
                    },
                    'package_status': package_status,
                    'validation_status': validation_status,
                    'attention_packages': attention_packages[:6],
                    'recent_activity': activities[:5],
                    'ai_health': {
                        'queue': Submission.objects.filter(status='ANALYZING').count(),
                        'success_rate': round((analyzed_count / (analyzed_count + failed_count) * 100), 1) if analyzed_count + failed_count else None,
                        'failures': failed_count,
                        'average_execution_ms': round(float(average_execution)) if average_execution is not None else None,
                    },
                    'alerts': alerts,
                },
            })

        if 'STUDENT' in role_set or user.is_superuser:
            mine = Submission.objects.filter(student=user)
            avg = LlmAnalysis.objects.filter(
                submission__student=user, is_current=True,
            ).aggregate(avg=Avg('percentage_correct'))['avg']
            summary.update({
                'my_submissions': mine.count(),
                'my_analyzed': mine.filter(status='ANALYZING').count(),
                'my_validated': mine.filter(status__in=['VALIDATED', 'PENDING_VALIDATION']).count(),
                'my_avg_score': round(float(avg), 2) if avg is not None else None,
            })

        if 'LECTURER' in role_set or user.is_superuser:
            in_scope = Q(subject_id__in=my_subjects_ids) if my_subjects_ids else Q()
            summary.update({
                'my_question_sets': QuestionSet.objects.filter(created_by=user).count(),
                'subject_question_sets': QuestionSet.objects.filter(in_scope).count(),
                'subject_submissions': Submission.objects.filter(in_scope).count(),
                'pending_validations': Submission.objects.filter(
                    in_scope & Q(status='PENDING_VALIDATION')
                ).count(),
            })

        return Response({
            'roles': roles,
            'is_superuser': user.is_superuser,
            'summary': summary,
        })


def require_admin(request):
    return request.user.is_superuser


HELP_ROLES = {choice for choice, _ in HelpArticle.Role.choices}


def serialize_help_article(article):
    return {'id': str(article.id), 'role': article.role, 'title': article.title, 'body': article.body, 'order_index': article.order_index, 'is_published': article.is_published}


class HelpArticleListView(APIView):
    authentication_classes = [TokenAuthentication]

    def get(self, request):
        role = request.query_params.get('role', 'GENERAL')
        if role not in HELP_ROLES:
            return Response({'detail': 'Role tidak valid.'}, status=status.HTTP_400_BAD_REQUEST)
        if not request.user.is_superuser:
            user_roles = set(UserRole.objects.filter(user=request.user).values_list('role', flat=True))
            if role not in user_roles:
                role = 'GENERAL'
        return Response([serialize_help_article(article) for article in HelpArticle.objects.filter(role=role, is_published=True)])


class AdminHelpArticleListView(APIView):
    authentication_classes = [TokenAuthentication]

    def get(self, request, role):
        if not require_admin(request): return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        if role not in HELP_ROLES: return Response({'detail': 'Role tidak valid.'}, status=status.HTTP_400_BAD_REQUEST)
        return Response([serialize_help_article(article) for article in HelpArticle.objects.filter(role=role)])

    def post(self, request, role):
        if not require_admin(request): return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        title, body = str(request.data.get('title', '')).strip(), str(request.data.get('body', '')).strip()
        if role not in HELP_ROLES or not title or not body: return Response({'detail': 'Role, judul, dan isi wajib diisi.'}, status=status.HTTP_400_BAD_REQUEST)
        article = HelpArticle.objects.create(role=role, title=title, body=body, order_index=int(request.data.get('order_index', 0)), is_published=bool(request.data.get('is_published', True)))
        return Response(serialize_help_article(article), status=status.HTTP_201_CREATED)


class AdminHelpArticleDetailView(APIView):
    authentication_classes = [TokenAuthentication]

    def patch(self, request, pk):
        if not require_admin(request): return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        try: article = HelpArticle.objects.get(pk=pk)
        except HelpArticle.DoesNotExist: return Response({'detail': 'Artikel tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)
        for field in ('title', 'body', 'order_index', 'is_published'):
            if field in request.data: setattr(article, field, request.data[field])
        if not str(article.title).strip() or not str(article.body).strip(): return Response({'detail': 'Judul dan isi wajib diisi.'}, status=status.HTTP_400_BAD_REQUEST)
        article.save(); return Response(serialize_help_article(article))

    def delete(self, request, pk):
        if not require_admin(request): return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        deleted, _ = HelpArticle.objects.filter(pk=pk).delete()
        return Response(status=status.HTTP_204_NO_CONTENT if deleted else status.HTTP_404_NOT_FOUND)


def serialize_website_section(section):
    return {
        'id': str(section.id), 'key': section.key, 'title': section.title,
        'eyebrow': section.eyebrow, 'body': section.body, 'image_url': section.image_url,
        'button_label': section.button_label, 'button_url': section.button_url,
        'content_json': section.content_json,
        'order_index': section.order_index, 'is_visible': section.is_visible,
        'is_system': section.is_system,
    }


class WebsiteSectionListView(APIView):
    authentication_classes = []

    def get(self, request):
        return Response([
            serialize_website_section(section) if section.is_visible else {'key': section.key, 'is_visible': False}
            for section in WebsiteSection.objects.all()
        ])


class AdminWebsiteSectionListView(APIView):
    authentication_classes = [TokenAuthentication]

    def get(self, request):
        if not require_admin(request): return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        return Response([serialize_website_section(section) for section in WebsiteSection.objects.all()])

    def post(self, request):
        if not require_admin(request): return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        title = str(request.data.get('title', '')).strip()
        key = slugify(str(request.data.get('key', '')).strip() or title)
        if not title or not key: return Response({'detail': 'Judul section wajib diisi.'}, status=status.HTTP_400_BAD_REQUEST)
        if WebsiteSection.objects.filter(key=key).exists(): return Response({'detail': 'Kunci section sudah digunakan.'}, status=status.HTTP_400_BAD_REQUEST)
        section = WebsiteSection.objects.create(
            key=key, title=title, eyebrow=str(request.data.get('eyebrow', '')).strip(),
            body=str(request.data.get('body', '')).strip(), image_url=str(request.data.get('image_url', '')).strip(),
            button_label=str(request.data.get('button_label', '')).strip(), button_url=str(request.data.get('button_url', '')).strip(),
            content_json=request.data.get('content_json', []),
            order_index=int(request.data.get('order_index', 0)), is_visible=bool(request.data.get('is_visible', True)),
        )
        return Response(serialize_website_section(section), status=status.HTTP_201_CREATED)


class AdminWebsiteSectionDetailView(APIView):
    authentication_classes = [TokenAuthentication]

    def patch(self, request, pk):
        if not require_admin(request): return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        try: section = WebsiteSection.objects.get(pk=pk)
        except WebsiteSection.DoesNotExist: return Response({'detail': 'Section tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)
        for field in ('title', 'eyebrow', 'body', 'image_url', 'button_label', 'button_url', 'content_json', 'order_index', 'is_visible'):
            if field in request.data: setattr(section, field, request.data[field])
        if not str(section.title).strip(): return Response({'detail': 'Judul section wajib diisi.'}, status=status.HTTP_400_BAD_REQUEST)
        section.save()
        return Response(serialize_website_section(section))

    def delete(self, request, pk):
        if not require_admin(request): return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        try: section = WebsiteSection.objects.get(pk=pk)
        except WebsiteSection.DoesNotExist: return Response(status=status.HTTP_404_NOT_FOUND)
        if section.is_system: return Response({'detail': 'Section bawaan dapat disembunyikan, tetapi tidak dapat dihapus.'}, status=status.HTTP_400_BAD_REQUEST)
        section.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


def managed_role_from_path(role):
    if role == 'lecturers':
        return UserRole.Role.LECTURER
    if role == 'students':
        return UserRole.Role.STUDENT
    return None


class AdminManagedUserListView(APIView):
    authentication_classes = [TokenAuthentication]

    def get(self, request, role):
        if not require_admin(request):
            return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        managed_role = managed_role_from_path(role)
        if managed_role is None:
            return Response({'detail': 'Role tidak valid.'}, status=status.HTTP_400_BAD_REQUEST)
        user_ids = UserRole.objects.filter(role=managed_role).values_list('user_id', flat=True)
        users = User.objects.filter(is_superuser=False, id__in=user_ids).order_by('full_name')
        return Response([self.serialize_user(user, managed_role) for user in users])

    def post(self, request, role):
        if not require_admin(request):
            return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        managed_role = managed_role_from_path(role)
        if managed_role is None:
            return Response({'detail': 'Role tidak valid.'}, status=status.HTTP_400_BAD_REQUEST)
        serializer = AdminManagedUserSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        if User.objects.filter(email__iexact=data['email']).exists():
            return Response({'email': ['Email sudah terdaftar.']}, status=status.HTTP_400_BAD_REQUEST)
        subject_ids = data.get('subject_ids', []) if managed_role == UserRole.Role.LECTURER else []
        valid_subjects = set(Subject.objects.filter(id__in=subject_ids).values_list('id', flat=True))
        if len(valid_subjects) != len(set(subject_ids)):
            return Response({'subject_ids': ['Ada mata kuliah yang tidak ditemukan.']}, status=status.HTTP_400_BAD_REQUEST)
        if not data.get('password'):
            return Response({'password': ['Password wajib diisi saat membuat akun.']}, status=status.HTTP_400_BAD_REQUEST)
        from django.contrib.auth.hashers import make_password
        user = User.objects.create(
            id=uuid.uuid4(), email=data['email'].lower(), full_name=data['full_name'],
            password_hash=make_password(data['password']), is_active=data.get('is_active', True),
        )
        UserRole.objects.create(user=user, role=managed_role)
        UserSubjectRole.objects.bulk_create([
            UserSubjectRole(user=user, subject_id=subject_id, role=managed_role)
            for subject_id in subject_ids
        ])
        return Response(self.serialize_user(user, managed_role), status=status.HTTP_201_CREATED)

    @staticmethod
    def serialize_user(user, role):
        roles = UserSubjectRole.objects.filter(
            user=user, role=UserSubjectRole.Role.LECTURER,
        ).select_related('subject') if role == UserRole.Role.LECTURER else []
        return {
            'id': str(user.id), 'email': user.email, 'full_name': user.full_name,
            'is_active': user.is_active, 'created_at': user.created_at,
            'role': role, 'subjects': [
                {'id': str(item.subject_id), 'name': item.subject.name, 'slug': item.subject.slug}
                for item in roles
            ],
        }


class AdminManagedUserDetailView(APIView):
    authentication_classes = [TokenAuthentication]

    def patch(self, request, role, pk):
        if not require_admin(request):
            return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        managed_role = managed_role_from_path(role)
        if managed_role is None:
            return Response({'detail': 'Role tidak valid.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            user = User.objects.get(pk=pk, is_superuser=False)
        except User.DoesNotExist:
            return Response({'detail': 'Akun tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)
        if not UserRole.objects.filter(user=user, role=managed_role).exists():
            return Response({'detail': 'Akun tidak termasuk role ini.'}, status=status.HTTP_404_NOT_FOUND)
        serializer = AdminManagedUserSerializer(data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        if 'email' in data and User.objects.exclude(pk=user.pk).filter(email__iexact=data['email']).exists():
            return Response({'email': ['Email sudah terdaftar.']}, status=status.HTTP_400_BAD_REQUEST)
        if managed_role == UserRole.Role.LECTURER and 'subject_ids' in data:
            subject_ids = data['subject_ids']
            valid_subjects = set(Subject.objects.filter(id__in=subject_ids).values_list('id', flat=True))
            if len(valid_subjects) != len(set(subject_ids)):
                return Response({'subject_ids': ['Ada mata kuliah yang tidak ditemukan.']}, status=status.HTTP_400_BAD_REQUEST)
        if 'email' in data: user.email = data['email'].lower()
        if 'full_name' in data: user.full_name = data['full_name']
        if 'is_active' in data: user.is_active = data['is_active']
        if data.get('password'):
            from django.contrib.auth.hashers import make_password
            user.password_hash = make_password(data['password'])
        user.save()
        if managed_role == UserRole.Role.LECTURER and 'subject_ids' in data:
            UserSubjectRole.objects.filter(user=user, role=managed_role).delete()
            UserSubjectRole.objects.bulk_create([
                UserSubjectRole(user=user, subject_id=subject_id, role=managed_role)
                for subject_id in subject_ids
            ])
        return Response(AdminManagedUserListView.serialize_user(user, managed_role))


class AdminSubjectListView(APIView):
    authentication_classes = [TokenAuthentication]

    def get(self, request):
        if require_admin(request):
            subjects = Subject.objects.order_by('name')
        else:
            subject_ids = UserSubjectRole.objects.filter(user=request.user, role=UserSubjectRole.Role.LECTURER).values_list('subject_id', flat=True)
            subjects = Subject.objects.filter(id__in=subject_ids).order_by('name')
        return Response([serialize_admin_subject(subject) for subject in subjects])

    def post(self, request):
        if not require_admin(request):
            return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        name = str(request.data.get('name', '')).strip()
        slug = slugify(str(request.data.get('slug', '')).strip() or name)
        if not name or not slug:
            return Response({'detail': 'Nama dan kode mata kuliah wajib diisi.'}, status=status.HTTP_400_BAD_REQUEST)
        if Subject.objects.filter(name__iexact=name).exists() or Subject.objects.filter(slug=slug).exists():
            return Response({'detail': 'Nama atau kode mata kuliah sudah digunakan.'}, status=status.HTTP_400_BAD_REQUEST)
        lecturer_ids, error = valid_lecturer_ids(request.data.get('lecturer_ids', []))
        if error:
            return Response({'lecturer_ids': [error]}, status=status.HTTP_400_BAD_REQUEST)
        subject = Subject.objects.create(
            id=uuid.uuid4(), name=name, slug=slug,
            description=str(request.data.get('description', '')).strip() or None,
            is_active=bool(request.data.get('is_active', True)),
        )
        replace_subject_lecturers(subject, lecturer_ids)
        return Response(serialize_admin_subject(subject), status=status.HTTP_201_CREATED)


def valid_lecturer_ids(raw_ids):
    lecturer_ids = list(dict.fromkeys(raw_ids or []))
    if not all(isinstance(item, str) for item in lecturer_ids):
        return [], 'Daftar dosen tidak valid.'
    found_ids = set(UserRole.objects.filter(
        user_id__in=lecturer_ids, role=UserRole.Role.LECTURER,
        user__is_superuser=False,
    ).values_list('user_id', flat=True))
    if len(found_ids) != len(lecturer_ids):
        return [], 'Ada dosen yang tidak ditemukan.'
    return lecturer_ids, None


def replace_subject_lecturers(subject, lecturer_ids):
    UserSubjectRole.objects.filter(subject=subject, role=UserSubjectRole.Role.LECTURER).delete()
    UserSubjectRole.objects.bulk_create([
        UserSubjectRole(user_id=lecturer_id, subject=subject, role=UserSubjectRole.Role.LECTURER)
        for lecturer_id in lecturer_ids
    ])


def serialize_admin_subject(subject):
    lecturer_roles = UserSubjectRole.objects.filter(
        subject=subject, role=UserSubjectRole.Role.LECTURER
    ).select_related('user').order_by('user__full_name')
    return {
        'id': str(subject.id), 'name': subject.name, 'slug': subject.slug,
        'description': subject.description or '', 'is_active': subject.is_active,
        'lecturers': [{'id': str(role.user_id), 'full_name': role.user.full_name, 'email': role.user.email} for role in lecturer_roles],
        'topic_count': Topic.objects.filter(subject=subject).count(),
    }


class AdminSubjectDetailView(APIView):
    authentication_classes = [TokenAuthentication]

    def get(self, request, pk):
        try:
            subject = Subject.objects.get(pk=pk)
        except Subject.DoesNotExist:
            return Response({'detail': 'Mata kuliah tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)
        if not require_admin(request) and not is_lecturer_for_subject(request.user, subject.id):
            return Response({'detail': 'Anda tidak memiliki akses ke mata kuliah ini.'}, status=status.HTTP_403_FORBIDDEN)
        return Response(serialize_admin_subject(subject))

    def patch(self, request, pk):
        if not require_admin(request):
            return Response({'detail': 'Akses administrator diperlukan.'}, status=status.HTTP_403_FORBIDDEN)
        try:
            subject = Subject.objects.get(pk=pk)
        except Subject.DoesNotExist:
            return Response({'detail': 'Mata kuliah tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)
        name = str(request.data.get('name', subject.name)).strip()
        slug = slugify(str(request.data.get('slug', subject.slug)).strip() or name)
        if not name or not slug:
            return Response({'detail': 'Nama dan kode mata kuliah wajib diisi.'}, status=status.HTTP_400_BAD_REQUEST)
        if Subject.objects.exclude(pk=subject.pk).filter(name__iexact=name).exists() or Subject.objects.exclude(pk=subject.pk).filter(slug=slug).exists():
            return Response({'detail': 'Nama atau kode mata kuliah sudah digunakan.'}, status=status.HTTP_400_BAD_REQUEST)
        if 'lecturer_ids' in request.data:
            lecturer_ids, error = valid_lecturer_ids(request.data['lecturer_ids'])
            if error:
                return Response({'lecturer_ids': [error]}, status=status.HTTP_400_BAD_REQUEST)
            replace_subject_lecturers(subject, lecturer_ids)
        subject.name = name
        subject.slug = slug
        if 'description' in request.data:
            subject.description = str(request.data['description']).strip() or None
        if 'is_active' in request.data:
            subject.is_active = bool(request.data['is_active'])
        subject.save()
        return Response(serialize_admin_subject(subject))


def serialize_admin_topic(topic):
    return {'id': str(topic.id), 'name': topic.name, 'description': topic.description or ''}


class AdminTopicListView(APIView):
    authentication_classes = [TokenAuthentication]

    def get(self, request, subject_id):
        if not require_admin(request) and not is_lecturer_for_subject(request.user, subject_id):
            return Response({'detail': 'Anda tidak memiliki akses ke mata kuliah ini.'}, status=status.HTTP_403_FORBIDDEN)
        return Response([serialize_admin_topic(topic) for topic in Topic.objects.filter(subject_id=subject_id).order_by('name')])

    def post(self, request, subject_id):
        try:
            subject = Subject.objects.get(pk=subject_id, is_active=True)
        except Subject.DoesNotExist:
            return Response({'detail': 'Mata kuliah aktif tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)
        if not require_admin(request) and not is_lecturer_for_subject(request.user, subject.id):
            return Response({'detail': 'Anda tidak memiliki akses ke mata kuliah ini.'}, status=status.HTTP_403_FORBIDDEN)
        name = str(request.data.get('name', '')).strip()
        if not name:
            return Response({'name': ['Nama topik wajib diisi.']}, status=status.HTTP_400_BAD_REQUEST)
        topic = Topic.objects.create(id=uuid.uuid4(), subject=subject, name=name, description=str(request.data.get('description', '')).strip() or None)
        return Response(serialize_admin_topic(topic), status=status.HTTP_201_CREATED)


class AdminTopicDetailView(APIView):
    authentication_classes = [TokenAuthentication]

    def patch(self, request, pk):
        try:
            topic = Topic.objects.get(pk=pk)
        except Topic.DoesNotExist:
            return Response({'detail': 'Topik tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)
        if not require_admin(request) and not is_lecturer_for_subject(request.user, topic.subject_id):
            return Response({'detail': 'Anda tidak memiliki akses ke topik ini.'}, status=status.HTTP_403_FORBIDDEN)
        name = str(request.data.get('name', topic.name)).strip()
        if not name:
            return Response({'name': ['Nama topik wajib diisi.']}, status=status.HTTP_400_BAD_REQUEST)
        topic.name = name
        if 'description' in request.data:
            topic.description = str(request.data['description']).strip() or None
        topic.save()
        return Response(serialize_admin_topic(topic))

    def delete(self, request, pk):
        try:
            topic = Topic.objects.get(pk=pk)
        except Topic.DoesNotExist:
            return Response({'detail': 'Topik tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)
        if not require_admin(request) and not is_lecturer_for_subject(request.user, topic.subject_id):
            return Response({'detail': 'Anda tidak memiliki akses ke topik ini.'}, status=status.HTTP_403_FORBIDDEN)
        topic.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


# ============================================================================
# SPRINT 2: QUESTION BANK VIEWS (LECTURER)
# ============================================================================

def is_lecturer_for_subject(user, subject_id):
    if user.is_superuser:
        return True
    return UserSubjectRole.objects.filter(
        user=user, subject_id=subject_id, role=UserSubjectRole.Role.LECTURER
    ).exists()


IMPORT_REQUIRED_COLUMNS = {'order_index', 'prompt', 'reference_answer'}


def _parse_import_indicators(raw_value):
    if not raw_value.strip():
        return [{'label': 'Konsep utama', 'description': '', 'weight': Decimal('1.0000')}]
    indicators = []
    for item in raw_value.split('|'):
        label, separator, weight = item.partition(':')
        if not separator:
            raise ValueError('Format indikator harus label:bobot|label:bobot.')
        try:
            parsed_weight = Decimal(weight.strip())
        except Exception as exc:
            raise ValueError(f'Bobot indikator tidak valid: {weight}.') from exc
        if not label.strip() or parsed_weight <= 0 or parsed_weight > 1:
            raise ValueError('Label indikator wajib diisi dan bobot harus antara 0 dan 1.')
        indicators.append({'label': label.strip(), 'description': '', 'weight': parsed_weight})
    if abs(sum(item['weight'] for item in indicators) - Decimal('1.0000')) > Decimal('0.0001'):
        raise ValueError('Total bobot indikator harus tepat 1.0000.')
    return indicators


def _normalize_import_row(raw):
    return {
        str(key).strip(): str(value).strip() if value is not None else ''
        for key, value in raw.items()
        if key
    }


def _read_xlsx_rows(upload, subject_name=None):
    if hasattr(upload, 'seek'):
        upload.seek(0)
    try:
        workbook = load_workbook(upload, read_only=False, data_only=True)
    except Exception as exc:
        raise ValueError('File Excel tidak dapat dibaca. Pastikan format file .xlsx valid.') from exc

    target_sheet = None
    if subject_name:
        for sname in workbook.sheetnames:
            if sname.strip().lower() == subject_name.strip().lower() or sname[:31].lower() == subject_name[:31].lower():
                target_sheet = workbook[sname]
                break

    if not target_sheet:
        for sname in workbook.sheetnames:
            ws = workbook[sname]
            for row in ws.iter_rows(values_only=True):
                row_str = [str(c).strip().upper() if c is not None else '' for c in row]
                if any(c in ('NOMOR_SOAL', 'ORDER_INDEX', 'NO_SOAL', 'PERTANYAAN_KONSEPTUAL', 'PROMPT') for c in row_str):
                    target_sheet = ws
                    break
            if target_sheet:
                break

    if not target_sheet:
        target_sheet = workbook.active

    header_row_idx = None
    headers = []
    for r_idx, row in enumerate(target_sheet.iter_rows(values_only=True), start=1):
        row_str = [str(c).strip().upper() if c is not None else '' for c in row]
        if any(c in ('NOMOR_SOAL', 'ORDER_INDEX', 'NO_SOAL', 'PERTANYAAN_KONSEPTUAL', 'PROMPT') for c in row_str):
            header_row_idx = r_idx
            headers = [str(c).strip() if c is not None else '' for c in row]
            break

    if not header_row_idx:
        raise ValueError('Baris header tidak ditemukan. Pastikan menggunakan template yang disediakan.')

    COLUMN_MAP = {
        'NOMOR_SOAL': 'order_index',
        'ORDER_INDEX': 'order_index',
        'NO_SOAL': 'order_index',
        'NO': 'order_index',
        'KODE_PAKET_UNIK': 'code',
        'KODE_PAKET': 'code',
        'CODE': 'code',
        'JUDUL_UJIAN': 'title',
        'JUDUL': 'title',
        'TITLE': 'title',
        'TOPIK': 'topic',
        'NAMA_TOPIK': 'topic',
        'TOPIC': 'topic',
        'DESKRIPSI_INSTRUKSI': 'description',
        'DESKRIPSI': 'description',
        'DESCRIPTION': 'description',
        'PERTANYAAN_KONSEPTUAL': 'prompt',
        'PERTANYAAN': 'prompt',
        'PROMPT': 'prompt',
        'JAWABAN_REFERENSI': 'reference_answer',
        'JAWABAN': 'reference_answer',
        'REFERENCE_ANSWER': 'reference_answer',
        'INDIKATOR_KONSEP': 'indicators',
        'INDICATORS': 'indicators',
        'ANSWER_KEY': 'answer_key',
    }

    normalized_headers = [COLUMN_MAP.get(h.upper(), h.lower()) for h in headers]

    rows = []
    for r_idx, row in enumerate(target_sheet.iter_rows(values_only=True), start=1):
        if r_idx <= header_row_idx:
            continue
        if not any(c is not None and str(c).strip() for c in row):
            continue
        raw_dict = dict(zip(normalized_headers, row))
        rows.append(_normalize_import_row({
            'order_index': raw_dict.get('order_index', len(rows) + 1),
            'prompt': raw_dict.get('prompt', ''),
            'reference_answer': raw_dict.get('reference_answer', ''),
            'indicators': raw_dict.get('indicators', '') or 'Ketepatan konsep:1.0000',
            'answer_key': raw_dict.get('answer_key', f"Q-{len(rows)+1}"),
            'code': raw_dict.get('code', ''),
            'title': raw_dict.get('title', ''),
            'topic': raw_dict.get('topic', ''),
            'description': raw_dict.get('description', ''),
        }))

    if not rows:
        raise ValueError(f'Sheet "{target_sheet.title}" tidak memiliki baris data soal.')
    return rows, target_sheet.title
    
def _read_question_import(upload, subject_name=None):
    if not upload:
        raise ValueError('File CSV atau Excel wajib diunggah.')
    if upload.size > 10 * 1024 * 1024:
        raise ValueError('Ukuran file maksimal 10 MB.')
    if upload.name.lower().endswith(('.xlsx', '.xlsm')):
        raw_rows, detected_sheet = _read_xlsx_rows(upload, subject_name=subject_name)
        headers = set(raw_rows[0]) if raw_rows else set()
    else:
        detected_sheet = None
        if hasattr(upload, 'seek'):
            upload.seek(0)
        try:
            text = upload.read().decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('File harus menggunakan encoding UTF-8.') from exc
        reader = csv.DictReader(io.StringIO(text))
        headers = {header.strip() for header in (reader.fieldnames or []) if header}
        raw_rows = list(reader)

    missing = sorted(IMPORT_REQUIRED_COLUMNS - headers)
    if missing:
        raise ValueError(f'Kolom wajib tidak ditemukan: {", ".join(missing)}.')

    rows = []
    for row_number, raw in enumerate(raw_rows, start=6):
        normalized = _normalize_import_row(raw)
        errors = []
        key = f"Q-{normalized.get('order_index', row_number - 5)}"
        if not normalized.get('prompt'):
            errors.append('prompt wajib diisi.')
        if not normalized.get('reference_answer'):
            errors.append('reference_answer wajib diisi.')
        try:
            order_index = int(float(normalized.get('order_index', '0')))
            if order_index < 1:
                order_index = len(rows) + 1
        except (ValueError, TypeError):
            order_index = len(rows) + 1

        try:
            indicators = _parse_import_indicators(normalized.get('indicators', ''))
        except ValueError as exc:
            indicators = []
            errors.append(str(exc))

        rows.append({
            'row_number': row_number,
            'question_key': key,
            'raw_data': normalized,
            'status': 'INVALID' if errors else 'VALID',
            'errors': errors,
            'order_index': order_index,
            'indicators': indicators,
        })

    return rows, detected_sheet


class QuestionImportCreateView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'detail': 'File wajib diunggah.'}, status=status.HTTP_400_BAD_REQUEST)

        subject = None
        subject_id = request.data.get('subject_id')
        if subject_id:
            subject = Subject.objects.filter(pk=subject_id).first()

        try:
            parsed_rows, detected_sheet = _read_question_import(file_obj, subject_name=subject.name if subject else None)
        except ValueError as exc:
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as exc:
            return Response({'detail': f'Terjadi error saat membaca format file: {str(exc)}'}, status=status.HTTP_400_BAD_REQUEST)

        if not subject and detected_sheet:
            subject = Subject.objects.filter(name__iexact=detected_sheet).first()

        if not subject:
            if request.user.is_superuser:
                subject = Subject.objects.filter(is_active=True).first()
            else:
                user_subjects = UserSubjectRole.objects.filter(
                    user=request.user, role=UserSubjectRole.Role.LECTURER
                ).values_list('subject_id', flat=True)
                subject = Subject.objects.filter(id__in=user_subjects, is_active=True).first()

        if not subject:
            return Response({'detail': 'Mata kuliah tidak ditemukan atau Anda belum memiliki akses ke mata kuliah terkait.'}, status=status.HTTP_400_BAD_REQUEST)

        # Validasi Topik dari Sheet
        distinct_topics = {
            r['raw_data'].get('topic', '').strip()
            for r in parsed_rows
            if r['raw_data'].get('topic', '').strip()
        }

        if len(distinct_topics) > 1:
            return Response({
                'detail': f"File impor hanya boleh berisi satu topik per paket ujian. Ditemukan beberapa topik: {', '.join(sorted(distinct_topics))}."
            }, status=status.HTTP_400_BAD_REQUEST)

        detected_topic_name = next(iter(distinct_topics)) if distinct_topics else ""
        if not detected_topic_name:
            return Response({
                'detail': "Kolom TOPIK wajib diisi untuk semua baris soal di dalam file Excel."
            }, status=status.HTTP_400_BAD_REQUEST)

        matched_topic = Topic.objects.filter(subject=subject, name__iexact=detected_topic_name).first()
        new_topic_detected = matched_topic is None

        if new_topic_detected and request.data.get('create_topic') in [True, 'true', '1']:
            matched_topic = Topic.objects.create(
                id=uuid.uuid4(),
                subject=subject,
                name=detected_topic_name,
                description=f"Dibuat otomatis dari impor bank soal.",
            )
            new_topic_detected = False

        first_data = parsed_rows[0]['raw_data'] if parsed_rows else {}
        code = (request.data.get('code') or first_data.get('code') or f"IMP-{uuid.uuid4().hex[:6].upper()}").strip().upper()
        title = (request.data.get('title') or first_data.get('title') or f"Paket Impor {subject.name}").strip()
        description = (request.data.get('description') or first_data.get('description') or '').strip()

        questions_payload = []
        for r in parsed_rows:
            raw = r['raw_data']
            if raw.get('prompt') and raw.get('reference_answer'):
                questions_payload.append({
                    'prompt': raw.get('prompt', ''),
                    'model_answer': raw.get('reference_answer', ''),
                    'indicators': [
                        {'label': 'Akurasi', 'description': 'Ketepatan konsep ilmiah dan kesesuaian prinsip dasar fisika.', 'weight': 40, 'isCustom': False},
                        {'label': 'Penjelasan', 'description': 'Kejelasan penalaran, alur argumen, dan langkah logika.', 'weight': 30, 'isCustom': False},
                        {'label': 'Kelengkapan', 'description': 'Kelengkapan seluruh variabel, satuan, dan elemen jawaban.', 'weight': 30, 'isCustom': False},
                    ]
                })

        if not questions_payload:
            return Response({'detail': 'Tidak ada butir soal yang valid di dalam file.'}, status=status.HTTP_400_BAD_REQUEST)

        return Response({
            'new_topic_detected': new_topic_detected,
            'detected_topic_name': detected_topic_name,
            'package': {
                'subject_id': str(subject.id),
                'topic_id': str(matched_topic.id) if matched_topic else None,
                'topic_name': matched_topic.name if matched_topic else detected_topic_name,
                'code': code,
                'title': title,
                'description': description,
                'questions': questions_payload,
            }
        }, status=status.HTTP_200_OK)


class QuestionImportTemplateView(APIView):
    authentication_classes = [TokenAuthentication]

    def get(self, request):
        subject_id = request.query_params.get('subject_id')
        target_subject = None
        if subject_id:
            target_subject = Subject.objects.filter(pk=subject_id, is_active=True).first()

        if not target_subject:
            target_subject = Subject.objects.filter(is_active=True).first()

        if not target_subject:
            target_subject = Subject(name="Physics", slug="physics")

        workbook = Workbook()
        ws = workbook.active
        sheet_title = target_subject.name[:31]
        ws.title = sheet_title

        banner_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        banner_font = Font(name="Calibri", size=11, bold=True, color="92400E")
        banner_border = Border(
            left=Side(style="thin", color="F59E0B"),
            right=Side(style="thin", color="F59E0B"),
            top=Side(style="thin", color="F59E0B"),
            bottom=Side(style="thin", color="F59E0B"),
        )

        header_fill = PatternFill(start_color="00288E", end_color="00288E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        thin_border = Border(
            left=Side(style="thin", color="E5E7EB"),
            right=Side(style="thin", color="E5E7EB"),
            top=Side(style="thin", color="E5E7EB"),
            bottom=Side(style="thin", color="E5E7EB"),
        )
        data_font = Font(name="Calibri", size=10)

        headers = [
            "NOMOR_SOAL",
            "KODE_PAKET_UNIK",
            "JUDUL_UJIAN",
            "TOPIK",
            "DESKRIPSI_INSTRUKSI",
            "PERTANYAAN_KONSEPTUAL",
            "JAWABAN_REFERENSI",
        ]

        ws.merge_cells("A1:G4")
        banner_cell = ws["A1"]
        banner_cell.value = (
            f"⚠️ SHEET MATA KULIAH: {target_subject.name.upper()}\n"
            f"Pastikan seluruh soal pada file ini diperuntukkan bagi mata kuliah {target_subject.name}.\n"
            f"Isi kolom pertanyaan dan jawaban referensi mulai dari baris ke-6."
        )
        banner_cell.fill = banner_fill
        banner_cell.font = banner_font
        banner_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for r in range(1, 5):
            for c in range(1, 8):
                ws.cell(row=r, column=c).border = banner_border
            ws.row_dimensions[r].height = 18

        ws.row_dimensions[5].height = 28
        for col_num, header in enumerate(headers, start=1):
            cell = ws.cell(row=5, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        sample_prefix = "FIS" if "phys" in target_subject.name.lower() else "BIO"
        sample_rows = [
            [
                1,
                f"{sample_prefix}-NEWT-01",
                f"Evaluasi Konseptual {target_subject.name} Bagian 1",
                "Hukum Newton",
                "Bacalah soal dengan saksama dan sertakan penalaran ilmiah.",
                "Mengapa berat semu seseorang di dalam lift yang dipercepat turun menjadi lebih kecil?",
                "Karena gaya normal N = m(g - a), percepatan lift mengurangi gaya kontak kaki pada timbangan.",
            ],
            [
                2,
                f"{sample_prefix}-NEWT-01",
                f"Evaluasi Konseptual {target_subject.name} Bagian 1",
                "Hukum Newton",
                "Bacalah soal dengan saksama dan sertakan penalaran ilmiah.",
                "Jelaskan mengapa gaya berat dan gaya normal pada balok diam bukan pasangan aksi-reaksi!",
                "Karena gaya normal dan gaya berat bekerja pada benda yang sama, sedangkan aksi-reaksi bekerja pada dua benda berbeda.",
            ],
        ]

        for row_idx, row_data in enumerate(sample_rows, start=6):
            ws.row_dimensions[row_idx].height = 36
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal="center" if col_idx in [1, 2] else "left",
                    vertical="center",
                    wrap_text=True,
                )

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.row < 5:
                    continue
                val_str = str(cell.value or "")
                lines = val_str.split("\n")
                longest_line = max(len(l) for l in lines) if lines else 0
                if longest_line > max_len:
                    max_len = longest_line
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 18), 55)

        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="template-{target_subject.slug}.xlsx"'
        return response
class QuestionExportView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            q_set = QuestionSet.objects.select_related('subject').get(pk=pk)
        except QuestionSet.DoesNotExist:
            return Response({'detail': 'Paket ujian tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)

        if not is_lecturer_for_subject(request.user, q_set.subject_id):
            return Response({'detail': 'Anda tidak berwenang mengekspor paket ini.'}, status=status.HTTP_403_FORBIDDEN)

        workbook = Workbook()
        ws = workbook.active
        ws.title = q_set.subject.name[:31]

        headers = [
            "NOMOR_SOAL",
            "KODE_PAKET_UNIK",
            "JUDUL_UJIAN",
            "DESKRIPSI_INSTRUKSI",
            "PERTANYAAN_KONSEPTUAL",
            "JAWABAN_REFERENSI",
        ]

        header_fill = PatternFill(start_color="00288E", end_color="00288E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for col_num, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        questions = Question.objects.filter(question_set=q_set).order_by('order_index')
        row_idx = 2
        for q in questions:
            v = QuestionVersion.objects.filter(question=q).order_by('-version_number').first()
            ws.cell(row=row_idx, column=1, value=q.order_index)
            ws.cell(row=row_idx, column=2, value=q_set.code)
            ws.cell(row=row_idx, column=3, value=q_set.title)
            ws.cell(row=row_idx, column=4, value=q_set.description or "")
            ws.cell(row=row_idx, column=5, value=v.prompt if v else "")
            ws.cell(row=row_idx, column=6, value=v.model_answer if v else "")
            row_idx += 1

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = 30

        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{q_set.code}-export.xlsx"'
        return response
        
class QuestionImportDetailView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            job = ImportJob.objects.select_related('subject').get(pk=pk, requested_by=request.user)
        except ImportJob.DoesNotExist:
            return Response({'detail': 'Import tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)
        return Response({
            'import_id': str(job.id), 'status': job.status, 'code': job.package_code,
            'title': job.package_title, 'total_rows': job.total_rows,
            'valid_rows': job.valid_rows, 'invalid_rows': job.invalid_rows,
            'errors': job.error_summary, 'rows': [
                {'row_number': row.row_number, 'question_key': row.question_key, 'status': row.status, 'errors': row.errors}
                for row in job.rows.order_by('row_number')
            ],
        })


class QuestionImportCommitView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            job = ImportJob.objects.get(pk=pk, requested_by=request.user)
        except ImportJob.DoesNotExist:
            return Response({'detail': 'Import tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)
        if job.status != 'READY_TO_IMPORT':
            return Response({'detail': 'Import belum valid dan tidak dapat di-commit.'}, status=status.HTTP_400_BAD_REQUEST)
        if QuestionSet.objects.filter(code__iexact=job.package_code).exists():
            return Response({'detail': 'Kode paket sudah digunakan.'}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            q_set = QuestionSet.objects.create(
                id=uuid.uuid4(), subject=job.subject, created_by=request.user,
                code=job.package_code, title=job.package_title,
                description=job.package_description or '', is_active=True,
            )
            for row in job.rows.order_by('row_number'):
                raw = row.raw_data
                question = Question.objects.create(
                    id=uuid.uuid4(), question_set=q_set,
                    external_key=row.question_key, order_index=int(raw['order_index']),
                )
                version = QuestionVersion.objects.create(
                    id=uuid.uuid4(), question=question, version_number=1,
                    prompt=raw['prompt'], model_answer=raw['reference_answer'],
                    is_published=False, created_by=request.user,
                )
                indicators = _parse_import_indicators(raw.get('indicators', ''))
                ConceptIndicator.objects.bulk_create([
                    ConceptIndicator(id=uuid.uuid4(), question_version=version, label=item['label'], description=item['description'], weight=item['weight'], order_index=index)
                    for index, item in enumerate(indicators, start=1)
                ])
                ReferenceAnswer.objects.create(
                    id=uuid.uuid4(), question_version=version,
                    answer_key=raw.get('answer_key') or row.question_key,
                    answer_text=raw['reference_answer'], answer_type='CANONICAL', is_primary=True,
                )
            job.status = 'IMPORTED'
            job.completed_at = timezone.now()
            job.save(update_fields=['status', 'completed_at'])
        return Response({'question_set_id': str(q_set.id), 'code': q_set.code, 'question_count': job.valid_rows, 'status': job.status})


class QuestionListCreateView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        subject_id = request.query_params.get('subject_id')
        topic_id = request.query_params.get('topic_id')

        if user.is_superuser:
            qs = QuestionSet.objects.all()
        else:
            lecturer_subjects = UserSubjectRole.objects.filter(
                user=user, role=UserSubjectRole.Role.LECTURER
            ).values_list('subject_id', flat=True)
            qs = QuestionSet.objects.filter(subject_id__in=lecturer_subjects)

        if subject_id:
            qs = qs.filter(subject_id=subject_id)
        if topic_id:
            qs = qs.filter(topic_id=topic_id)

        qs = qs.select_related('subject', 'topic', 'created_by').order_by('-created_at')
        set_ids = [item.id for item in qs]

        submission_stats = {}
        if set_ids:
            with connection.cursor() as cursor:
                cursor.execute(
                    """
                    SELECT 
                        q.question_set_id,
                        COUNT(s.id) AS total_submissions,
                        COUNT(DISTINCT s.student_id) AS distinct_students,
                        COUNT(s.id) FILTER (WHERE s.status = 'PENDING_VALIDATION') AS pending_validations
                    FROM questions q
                    JOIN question_versions qv ON qv.question_id = q.id
                    JOIN submissions s ON s.question_version_id = qv.id
                    WHERE q.question_set_id = ANY(%s)
                    GROUP BY q.question_set_id
                    """,
                    [set_ids],
                )
                for row in cursor.fetchall():
                    submission_stats[row[0]] = {
                        'total_submissions': row[1],
                        'distinct_students': row[2],
                        'pending_validations': row[3],
                    }

        results = []
        for item in qs:
            set_questions = list(Question.objects.filter(question_set=item).order_by('order_index'))
            version_data = []
            for question in set_questions:
                v = QuestionVersion.objects.filter(question=question).order_by('-version_number').first()
                if v:
                    version_data.append({
                        'id': str(v.id),
                        'question_id': str(question.id),
                        'order_index': question.order_index,
                        'version_number': v.version_number,
                        'prompt_preview': v.prompt[:160] + ('...' if len(v.prompt) > 160 else ''),
                        'is_published': v.is_published,
                    })

            stats = submission_stats.get(item.id, {
                'total_submissions': 0,
                'distinct_students': 0,
                'pending_validations': 0,
            })

            results.append({
                'id': str(item.id),
                'code': item.code,
                'title': item.title,
                'description': item.description or '',
                'subject_id': str(item.subject_id),
                'subject_name': item.subject.name,
                'topic_id': str(item.topic_id) if item.topic_id else None,
                'topic_name': item.topic.name if item.topic else None,
                'is_active': item.is_active,
                'created_by_name': item.created_by.full_name,
                'created_at': item.created_at,
                'question_count': len(set_questions),
                'total_submissions_count': stats['total_submissions'],
                'distinct_students_count': stats['distinct_students'],
                'pending_validations_count': stats['pending_validations'],
                'latest_versions': version_data,
                'latest_version': version_data[0] if version_data else None,
            })

        return Response(results)

    def post(self, request):
        serializer = QuestionSetCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        subject_id = data['subject_id']
        if not is_lecturer_for_subject(request.user, subject_id):
            return Response(
                {'detail': 'Anda tidak berwenang menambahkan soal pada mata kuliah ini.'},
                status=status.HTTP_403_FORBIDDEN
            )

        code = data['code']
        if QuestionSet.objects.filter(code__iexact=code).exists():
            return Response(
                {'code': ['Kode soal sudah digunakan. Gunakan kode unik lainnya.']},
                status=status.HTTP_400_BAD_REQUEST
            )

        with transaction.atomic():
            q_set = QuestionSet.objects.create(
                id=uuid.uuid4(),
                subject_id=subject_id,
                topic_id=data.get('topic_id'),
                created_by=request.user,
                code=code,
                title=data['title'],
                description=data.get('description', ''),
                is_active=True,
            )

            questions_list = data.get('questions', [])
            if not questions_list:
                questions_list = [{
                    'prompt': data['prompt'],
                    'model_answer': data['model_answer'],
                    'indicators': data.get('indicators', []),
                }]

            for q_idx, q_item in enumerate(questions_list, start=1):
                q = Question.objects.create(
                    id=uuid.uuid4(),
                    question_set=q_set,
                    order_index=q_idx,
                )

                qv = QuestionVersion.objects.create(
                    id=uuid.uuid4(),
                    question=q,
                    version_number=1,
                    prompt=q_item['prompt'],
                    model_answer=q_item['model_answer'],
                    is_published=False,
                    created_by=request.user,
                )

                ReferenceAnswer.objects.create(
                    id=uuid.uuid4(),
                    question_version=qv,
                    answer_key=f"ANS-{q_idx:03d}",
                    answer_text=q_item['model_answer'],
                    answer_type='CANONICAL',
                    is_primary=True,
                )

                indicators_data = q_item.get('indicators', [])
                for idx, ind in enumerate(indicators_data, start=1):
                    ConceptIndicator.objects.create(
                        id=uuid.uuid4(),
                        question_version=qv,
                        label=ind['label'],
                        description=ind.get('description', ''),
                        weight=ind['weight'],
                        order_index=idx,
                    )

                if data.get('publish', False):
                    with connection.cursor() as cursor:
                        cursor.execute("CALL sp_publish_question_version(%s, %s);", [str(qv.id), str(request.user.id)])
                    qv.refresh_from_db()

        return Response({
            'id': str(q_set.id),
            'code': q_set.code,
            'title': q_set.title,
            'version_id': str(qv.id),
            'is_published': qv.is_published,
            'is_active': q_set.is_active,
            'message': 'Soal berhasil disimpan.',
        }, status=status.HTTP_201_CREATED)


class QuestionDetailView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            q_set = QuestionSet.objects.select_related('subject', 'topic', 'created_by').get(pk=pk)
        except QuestionSet.DoesNotExist:
            return Response({'detail': 'Soal tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)

        if not is_lecturer_for_subject(request.user, q_set.subject_id):
            return Response({'detail': 'Anda tidak berwenang mengakses soal ini.'}, status=status.HTTP_403_FORBIDDEN)

        questions = []
        for question in Question.objects.filter(question_set=q_set).order_by('order_index'):
            versions = []
            for version in QuestionVersion.objects.filter(question=question).order_by('-version_number'):
                indicators = ConceptIndicator.objects.filter(question_version=version).order_by('order_index')

                refs = list(
                    ReferenceAnswer.objects.filter(question_version=version).order_by('-is_primary', 'created_at')
                )

                if refs:
                    ref_answers = [
                        {
                            'id': str(reference.id),
                            'answer_key': reference.answer_key,
                            'answer_text': reference.answer_text,
                            'answer_type': reference.answer_type,
                            'is_primary': reference.is_primary,
                        }
                        for reference in refs
                    ]
                else:
                    ref_answers = [
                        {
                            'id': str(version.id),
                            'answer_key': 'CANONICAL',
                            'answer_text': version.model_answer,
                            'answer_type': 'CANONICAL',
                            'is_primary': True,
                        }
                    ]

                versions.append({
                    'id': str(version.id),
                    'version_number': version.version_number,
                    'prompt': version.prompt,
                    'model_answer': version.model_answer,
                    'is_published': version.is_published,
                    'created_at': version.created_at,
                    'reference_answers': ref_answers,
                    'indicators': [
                        {
                            'id': str(indicator.id),
                            'label': indicator.label,
                            'description': indicator.description or '',
                            'weight': str(indicator.weight),
                            'order_index': indicator.order_index,
                        }
                        for indicator in indicators
                    ],
                })
            questions.append({
                'id': str(question.id),
                'external_key': question.external_key,
                'order_index': question.order_index,
                'versions': versions,
            })

        return Response({
            'id': str(q_set.id),
            'code': q_set.code,
            'title': q_set.title,
            'description': q_set.description or '',
            'subject_id': str(q_set.subject_id),
            'subject_name': q_set.subject.name,
            'topic_id': str(q_set.topic_id) if q_set.topic_id else None,
            'topic_name': q_set.topic.name if q_set.topic else None,
            'is_active': q_set.is_active,
            'created_at': q_set.created_at,
            'questions': questions,
        })

    def put(self, request, pk):
        try:
            q_set = QuestionSet.objects.get(pk=pk)
        except QuestionSet.DoesNotExist:
            return Response({'detail': 'Soal tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)

        if not is_lecturer_for_subject(request.user, q_set.subject_id):
            return Response({'detail': 'Anda tidak berwenang mengubah soal ini.'}, status=status.HTTP_403_FORBIDDEN)

        serializer = QuestionSetUpdateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        with transaction.atomic():
            if 'title' in data:
                q_set.title = data['title']
            if 'description' in data:
                q_set.description = data['description']
            if 'topic_id' in data:
                q_set.topic_id = data['topic_id']
            q_set.save()

            q = Question.objects.filter(question_set=q_set).order_by('order_index').first()
            if not q:
                q = Question.objects.create(id=uuid.uuid4(), question_set=q_set, order_index=1)

            latest_v = QuestionVersion.objects.filter(question=q).order_by('-version_number').first()

            prompt = data.get('prompt', latest_v.prompt if latest_v else "")
            model_answer = data.get('model_answer', latest_v.model_answer if latest_v else "")

            if latest_v and latest_v.is_published:
                new_v_id = uuid.uuid4()
                next_version = latest_v.version_number + 1
                target_v = QuestionVersion.objects.create(
                    id=new_v_id,
                    question=q,
                    version_number=next_version,
                    prompt=prompt,
                    model_answer=model_answer,
                    is_published=False,
                    created_by=request.user,
                )
                ReferenceAnswer.objects.create(
                    id=uuid.uuid4(),
                    question_version=target_v,
                    answer_key='CANONICAL',
                    answer_text=model_answer,
                    answer_type='CANONICAL',
                    is_primary=True,
                )
            elif latest_v:
                latest_v.prompt = prompt
                latest_v.model_answer = model_answer
                latest_v.save()
                target_v = latest_v
                ref = ReferenceAnswer.objects.filter(question_version=target_v, is_primary=True).first()
                if ref:
                    ref.answer_text = model_answer
                    ref.save(update_fields=['answer_text'])
                else:
                    ReferenceAnswer.objects.create(
                        id=uuid.uuid4(),
                        question_version=target_v,
                        answer_key='CANONICAL',
                        answer_text=model_answer,
                        answer_type='CANONICAL',
                        is_primary=True,
                    )
            else:
                target_v = QuestionVersion.objects.create(
                    id=uuid.uuid4(),
                    question=q,
                    version_number=1,
                    prompt=prompt,
                    model_answer=model_answer,
                    is_published=False,
                    created_by=request.user,
                )
                ReferenceAnswer.objects.create(
                    id=uuid.uuid4(),
                    question_version=target_v,
                    answer_key='CANONICAL',
                    answer_text=model_answer,
                    answer_type='CANONICAL',
                    is_primary=True,
                )

            if 'indicators' in data:
                ConceptIndicator.objects.filter(question_version=target_v).delete()
                for idx, ind in enumerate(data['indicators'], start=1):
                    ConceptIndicator.objects.create(
                        id=uuid.uuid4(),
                        question_version=target_v,
                        label=ind['label'],
                        description=ind.get('description', ''),
                        weight=ind['weight'],
                        order_index=idx,
                    )

            if data.get('publish', False) and not target_v.is_published:
                with connection.cursor() as cursor:
                    cursor.execute("CALL sp_publish_question_version(%s, %s);", [str(target_v.id), str(request.user.id)])
                target_v.refresh_from_db()

        return Response({'detail': 'Soal berhasil diperbarui.'})




class QuestionSetReviewView(APIView):
    """Lecturer progress for students who submitted this package."""

    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            q_set = QuestionSet.objects.select_related('subject').get(pk=pk)
        except QuestionSet.DoesNotExist:
            return Response({'detail': 'Paket ujian tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)

        if not is_lecturer_for_subject(request.user, q_set.subject_id):
            return Response({'detail': 'Anda tidak berwenang mengakses paket ini.'}, status=status.HTTP_403_FORBIDDEN)

        questions = list(Question.objects.filter(question_set=q_set).order_by('order_index'))
        question_ids = {question.id for question in questions}
        versions = list(QuestionVersion.objects.filter(question_id__in=question_ids))
        versions_by_id = {version.id: version for version in versions}
        published_question_ids = {
            version.question_id for version in versions if version.is_published
        }

        submissions = list(
            Submission.objects.filter(
                question_version_id__in=versions_by_id,
            ).order_by('-submitted_at', '-attempt_no')
        )
        student_ids = {submission.student_id for submission in submissions}
        students = list(User.objects.filter(
            id__in=student_ids, is_active=True,
            userrole__role=UserRole.Role.STUDENT,
        ).order_by('full_name', 'email').distinct())
        student_ids = {student.id for student in students}
        submissions = [submission for submission in submissions if submission.student_id in student_ids]

        analyses = {
            analysis.submission_id: analysis
            for analysis in LlmAnalysis.objects.filter(
                submission_id__in=[submission.id for submission in submissions],
                is_current=True,
            )
        }
        validations = {
            validation.analysis_id: validation
            for validation in Validation.objects.filter(
                analysis_id__in=[analysis.id for analysis in analyses.values()]
            )
        }

        student_rows = []
        for student in students:
            student_subs = [s for s in submissions if s.student_id == student.id]
            answered_questions = {
                versions_by_id[s.question_version_id].question_id
                for s in student_subs if s.question_version_id in versions_by_id
            }

            all_submissions_payload = []
            for s in student_subs:
                version = versions_by_id.get(s.question_version_id)
                if not version:
                    continue
                q = next((item for item in questions if item.id == version.question_id), None)
                analysis = analyses.get(s.id)
                validation = validations.get(analysis.id) if analysis else None

                all_submissions_payload.append({
                    'question_id': str(version.question_id),
                    'order_index': q.order_index if q else 1,
                    'question_prompt_preview': (version.prompt[:160] + ('...' if len(version.prompt) > 160 else '')),
                    'submission_id': str(s.id),
                    'attempt_no': s.attempt_no,
                    'status': s.status,
                    'submitted_at': s.submitted_at,
                    'analysis_id': str(analysis.id) if analysis else None,
                    'validation_status': validation.status if validation else None,
                })

            student_rows.append({
                'student_id': str(student.id),
                'student_name': student.full_name,
                'student_email': student.email,
                'answered_count': len(answered_questions),
                'published_question_count': len(published_question_ids),
                'total_attempts_count': len(student_subs),
                'all_submissions': all_submissions_payload,
            })

        return Response({
            'id': str(q_set.id),
            'code': q_set.code,
            'title': q_set.title,
            'description': q_set.description or '',
            'subject_id': str(q_set.subject_id),
            'subject_name': q_set.subject.name,
            'is_active': q_set.is_active,
            'published_question_count': len(published_question_ids),
            'roster_scope': 'submitted_students',
            'unsubmitted_roster_available': False,
            'students': student_rows,
        })


class QuestionSetStudentReviewView(APIView):
    """Full package review for one student, supporting all submission attempts."""

    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request, pk, student_id):
        try:
            q_set = QuestionSet.objects.select_related('subject').get(pk=pk)
        except QuestionSet.DoesNotExist:
            return Response({'detail': 'Paket ujian tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)

        if not is_lecturer_for_subject(request.user, q_set.subject_id):
            return Response({'detail': 'Anda tidak berwenang mengakses paket ini.'}, status=status.HTTP_403_FORBIDDEN)

        try:
            student = User.objects.filter(
                pk=student_id, is_active=True, userrole__role=UserRole.Role.STUDENT,
                submissions__question_version_id__in=QuestionVersion.objects.filter(
                    question__question_set=q_set,
                ).values('id'),
            ).distinct().get()
        except User.DoesNotExist:
            return Response({'detail': 'Mahasiswa belum mengumpulkan paket ini.'}, status=status.HTTP_404_NOT_FOUND)

        questions = list(Question.objects.filter(question_set=q_set).order_by('order_index'))
        question_ids = [question.id for question in questions]
        published_versions = list(
            QuestionVersion.objects.filter(question_id__in=question_ids, is_published=True)
            .order_by('question_id', '-version_number')
        )
        version_by_question = {}
        for version in published_versions:
            version_by_question.setdefault(version.question_id, version)

        package_versions = list(QuestionVersion.objects.filter(question_id__in=question_ids))
        package_version_ids = [version.id for version in package_versions]
        
        # Fetch ALL submissions from this student for this package
        submissions = list(
            Submission.objects.filter(
                student=student,
                question_version_id__in=package_version_ids,
            ).order_by('-submitted_at', '-attempt_no')
        )
        version_question_ids = {version.id: version.question_id for version in package_versions}

        analyses = {
            analysis.submission_id: analysis
            for analysis in LlmAnalysis.objects.filter(
                submission_id__in=[submission.id for submission in submissions],
                is_current=True,
            )
        }
        validations = {
            validation.analysis_id: validation
            for validation in Validation.objects.filter(
                analysis_id__in=[analysis.id for analysis in analyses.values()]
            ).select_related('lecturer')
        }

        indicators_by_version = {}
        for indicator in ConceptIndicator.objects.filter(
            question_version_id__in=[v.id for v in version_by_question.values()]
        ).order_by('order_index'):
            indicators_by_version.setdefault(indicator.question_version_id, []).append({
                'id': str(indicator.id), 'label': indicator.label,
                'description': indicator.description or '', 'weight': str(indicator.weight),
                'order_index': indicator.order_index,
            })

        rows = []
        for question in questions:
            version = version_by_question.get(question.id)
            if not version:
                continue

            # Group all attempts for this question
            question_subs = [
                s for s in submissions if version_question_ids.get(s.question_version_id) == question.id
            ]

            attempts_payload = []
            for s in question_subs:
                analysis = analyses.get(s.id)
                validation = validations.get(analysis.id) if analysis else None

                attempts_payload.append({
                    'submission_id': str(s.id),
                    'attempt_no': s.attempt_no,
                    'status': s.status,
                    'answer_text': s.answer_text,
                    'submitted_at': s.submitted_at,
                    'analysis': {
                        'id': str(analysis.id),
                        'run_number': analysis.run_number,
                        'percentage_correct': str(analysis.percentage_correct),
                        'tier_level': analysis.tier_level_snapshot,
                        'tier_label': analysis.tier_label_snapshot,
                        'confidence': str(analysis.confidence),
                        'explanation': analysis.explanation,
                        'execution_time_ms': analysis.execution_time_ms,
                        'concept_breakdown_json': analysis.concept_breakdown_json or {},
                        'validation': {
                            'status': validation.status,
                            'final_percentage': str(validation.final_percentage) if validation.final_percentage is not None else None,
                            'final_tier_level': validation.final_tier_level_snapshot,
                            'final_feedback': validation.final_feedback,
                            'lecturer_name': validation.lecturer.full_name,
                            'validated_at': validation.validated_at,
                        } if validation else None,
                    } if analysis else None,
                })

            attempts_payload.sort(key=lambda a: a['attempt_no'], reverse=True)

            rows.append({
                'question_id': str(question.id),
                'order_index': question.order_index,
                'version_id': str(version.id),
                'version_number': version.version_number,
                'prompt': version.prompt,
                'model_answer': version.model_answer,
                'indicators': indicators_by_version.get(version.id, []),
                'status': attempts_payload[0]['status'] if attempts_payload else 'UNANSWERED',
                'attempts': attempts_payload,
            })

        return Response({
            'package': {
                'id': str(q_set.id), 'code': q_set.code, 'title': q_set.title,
                'description': q_set.description or '', 'subject_id': str(q_set.subject_id),
                'subject_name': q_set.subject.name,
            },
            'student': {'id': str(student.id), 'name': student.full_name, 'email': student.email},
            'published_question_count': len(rows),
            'answered_count': sum(len(row['attempts']) > 0 for row in rows),
            'total_attempts_count': sum(len(row['attempts']) for row in rows),
            'questions': rows,
        })


class QuestionToggleActiveView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def patch(self, request, pk):
        try:
            q_set = QuestionSet.objects.get(pk=pk)
        except QuestionSet.DoesNotExist:
            return Response({'detail': 'Soal tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)

        if not is_lecturer_for_subject(request.user, q_set.subject_id):
            return Response({'detail': 'Anda tidak berwenang mengelola status soal ini.'}, status=status.HTTP_403_FORBIDDEN)

        q_set.is_active = not q_set.is_active
        q_set.save(update_fields=['is_active', 'updated_at'])

        status_text = "diaktifkan" if q_set.is_active else "dinonaktifkan"
        return Response({
            'id': str(q_set.id),
            'is_active': q_set.is_active,
            'message': f'Soal berhasil {status_text}.'
        })


class QuestionSetPublishView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            q_set = QuestionSet.objects.get(pk=pk)
        except QuestionSet.DoesNotExist:
            return Response({'detail': 'Paket ujian tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)

        if not is_lecturer_for_subject(request.user, q_set.subject_id):
            return Response({'detail': 'Anda tidak berwenang menerbitkan paket ini.'}, status=status.HTTP_403_FORBIDDEN)

        questions = list(Question.objects.filter(question_set=q_set).order_by('order_index'))
        if not questions:
            return Response({'detail': 'Paket harus memiliki minimal satu pertanyaan.'}, status=status.HTTP_400_BAD_REQUEST)

        versions = []
        for question in questions:
            version = QuestionVersion.objects.filter(question=question).order_by('-version_number').first()
            if not version:
                return Response({'detail': f'Pertanyaan ke-{question.order_index} belum memiliki versi.'}, status=status.HTTP_400_BAD_REQUEST)
            total_weight = ConceptIndicator.objects.filter(question_version=version).aggregate(total=Sum('weight'))['total'] or Decimal('0')
            if abs(total_weight - Decimal('1.0000')) > Decimal('0.0001'):
                return Response({'detail': f'Bobot indikator pertanyaan ke-{question.order_index} harus tepat 1.0000.'}, status=status.HTTP_400_BAD_REQUEST)
            versions.append(version)

        try:
            with transaction.atomic():
                for version in versions:
                    if not version.is_published:
                        with connection.cursor() as cursor:
                            cursor.execute("CALL sp_publish_question_version(%s, %s);", [str(version.id), str(request.user.id)])
        except Exception as exc:
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        return Response({
            'id': str(q_set.id),
            'code': q_set.code,
            'question_count': len(versions),
            'is_published': True,
            'message': 'Paket ujian berhasil diterbitkan.',
        })


class QuestionPublishView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def post(self, request, version_id):
        try:
            qv = QuestionVersion.objects.select_related('question__question_set').get(pk=version_id)
        except QuestionVersion.DoesNotExist:
            return Response({'detail': 'Versi soal tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)

        subject_id = qv.question.question_set.subject_id
        if not is_lecturer_for_subject(request.user, subject_id):
            return Response({'detail': 'Anda tidak berwenang mempublikasikan soal ini.'}, status=status.HTTP_403_FORBIDDEN)

        try:
            with connection.cursor() as cursor:
                cursor.execute("CALL sp_publish_question_version(%s, %s);", [str(qv.id), str(request.user.id)])
            qv.refresh_from_db()
            return Response({'detail': 'Versi soal berhasil dipublikasikan.', 'is_published': qv.is_published})
        except Exception as e:
            return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)


# SPRINT 3: STUDENT SUBMISSION API (UC-01 / P3)
# ============================================================================

def is_active_student(user):
    if user.is_superuser:
        return True
    return user.is_active and UserRole.objects.filter(
        user=user, role=UserRole.Role.STUDENT,
    ).exists()


class StudentSetLookupView(APIView):
    """Lookup soal set by kode (UC-01: mahasiswa memasukkan kode soal)."""

    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request):
        code = (request.query_params.get('code') or '').strip()
        if not code:
            return Response(
                {'detail': 'Parameter kode wajib diisi.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            q_set = QuestionSet.objects.select_related('subject', 'created_by').get(
                code__iexact=code, is_active=True
            )
        except QuestionSet.DoesNotExist:
            return Response(
                {'detail': 'Kode soal tidak ditemukan.'},
                status=status.HTTP_404_NOT_FOUND,
            )

        if not is_active_student(request.user):
            return Response(
                {'detail': 'Akun mahasiswa aktif diperlukan.'},
                status=status.HTTP_403_FORBIDDEN,
            )

        questions = Question.objects.filter(question_set=q_set).order_by('order_index')

        items = []
        for q in questions:
            version = QuestionVersion.objects.filter(
                question=q, is_published=True
            ).order_by('-version_number').first()
            if not version:
                continue
            latest_sub = Submission.objects.filter(
                student=request.user, question_version_id=version.id
            ).order_by('-attempt_no').first()
            items.append({
                'question_id': str(q.id),
                'order_index': q.order_index,
                'version_id': str(version.id),
                'version_number': version.version_number,
                'prompt': version.prompt,
                'latest_submission': {
                    'submission_id': str(latest_sub.id),
                    'attempt_no': latest_sub.attempt_no,
                    'status': latest_sub.status,
                    'submitted_at': latest_sub.submitted_at,
                } if latest_sub else None,
            })

        return Response({
            'id': str(q_set.id),
            'code': q_set.code,
            'title': q_set.title,
            'description': q_set.description or '',
            'subject_id': str(q_set.subject_id),
            'subject_name': q_set.subject.name,
            'is_active': q_set.is_active,
            'questions': items,
        })


class StudentSubmissionCreateView(APIView):
    """Simpan jawaban konseptual mahasiswa (UC-01; sp_submit_conceptual_answer)."""

    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def post(self, request, pk, qid):
        try:
            q_set = QuestionSet.objects.get(pk=pk, is_active=True)
        except QuestionSet.DoesNotExist:
            return Response(
                {'detail': 'Soal tidak ditemukan.'},
                status=status.HTTP_404_NOT_FOUND,
            )

        try:
            question = Question.objects.get(pk=qid, question_set=q_set)
        except Question.DoesNotExist:
            return Response(
                {'detail': 'Pertanyaan tidak ditemukan.'},
                status=status.HTTP_404_NOT_FOUND,
            )

        version = QuestionVersion.objects.filter(
            question=question, is_published=True
        ).order_by('-version_number').first()
        if not version:
            return Response(
                {'detail': 'Versi soal belum dipublikasikan. Tidak dapat mengumpulkan jawaban.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not is_active_student(request.user):
            return Response(
                {'detail': 'Akun mahasiswa aktif diperlukan.'},
                status=status.HTTP_403_FORBIDDEN,
            )

        serializer = SubmissionCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        answer_text = serializer.validated_data['answer_text']

        try:
            with transaction.atomic():
                with connection.cursor() as cursor:
                    cursor.execute(
                        "CALL sp_submit_conceptual_answer(%s, %s, %s, NULL);",
                        [str(request.user.id), str(version.id), answer_text],
                    )
                    row = cursor.fetchone()
                    submission_id = str(row[0]) if row and row[0] else None
                    if not submission_id:
                        cursor.execute(
                            """
                            SELECT id FROM submissions
                            WHERE student_id = %s AND question_version_id = %s
                            ORDER BY attempt_no DESC LIMIT 1
                            """,
                            [str(request.user.id), str(version.id)],
                        )
                        fallback = cursor.fetchone()
                        submission_id = str(fallback[0]) if fallback else None
        except Exception as e:
            return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)

        if not submission_id:
            return Response(
                {'detail': 'Gagal menyimpan jawaban. Silakan coba lagi.'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        submission = Submission.objects.get(pk=submission_id)
        return Response({
            'submission_id': str(submission.id),
            'question_id': str(question.id),
            'question_version_id': str(version.id),
            'attempt_no': submission.attempt_no,
            'status': submission.status,
            'submitted_at': submission.submitted_at,
        }, status=status.HTTP_201_CREATED)


class StudentPackageSubmissionCreateView(APIView):
    """Submit every published question in a set as one atomic package."""

    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        serializer = PackageSubmissionCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        answers_by_question = {
            answer['question_id']: answer['answer_text']
            for answer in serializer.validated_data['answers']
        }

        try:
            with transaction.atomic():
                try:
                    q_set = QuestionSet.objects.select_for_update().get(pk=pk, is_active=True)
                except QuestionSet.DoesNotExist:
                    return Response(
                        {'detail': 'Soal tidak ditemukan.'},
                        status=status.HTTP_404_NOT_FOUND,
                    )

                if not is_active_student(request.user):
                    return Response(
                        {'detail': 'Akun mahasiswa aktif diperlukan.'},
                        status=status.HTTP_403_FORBIDDEN,
                    )

                questions = list(
                    Question.objects.select_for_update().filter(question_set=q_set).order_by('order_index')
                )
                published = []
                for question in questions:
                    version = (
                        QuestionVersion.objects.select_for_update()
                        .filter(question=question, is_published=True)
                        .order_by('-version_number')
                        .first()
                    )
                    if version:
                        published.append((question, version))

                published_question_ids = {question.id for question, _ in published}
                if not published:
                    return Response(
                        {'detail': 'Belum ada pertanyaan yang dipublikasikan pada bank soal ini.'},
                        status=status.HTTP_400_BAD_REQUEST,
                    )
                if set(answers_by_question) != published_question_ids:
                    return Response(
                        {'detail': 'Semua pertanyaan yang dipublikasikan harus dijawab tepat satu kali.'},
                        status=status.HTTP_400_BAD_REQUEST,
                    )

                submission_ids = []
                for question, version in published:
                    with connection.cursor() as cursor:
                        cursor.execute(
                            "CALL sp_submit_conceptual_answer(%s, %s, %s, NULL);",
                            [str(request.user.id), str(version.id), answers_by_question[question.id]],
                        )
                        row = cursor.fetchone()
                        submission_id = str(row[0]) if row and row[0] else None
                        if not submission_id:
                            raise RuntimeError('Gagal menyimpan salah satu jawaban.')
                    submission_ids.append((question, version, submission_id))
        except Exception as exc:
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        submissions = {
            str(submission.id): submission
            for submission in Submission.objects.filter(
                id__in=[submission_id for _, _, submission_id in submission_ids]
            )
        }
        return Response({
            'set_id': str(q_set.id),
            'submissions': [
                {
                    'submission_id': submission_id,
                    'question_id': str(question.id),
                    'question_version_id': str(version.id),
                    'attempt_no': submissions[submission_id].attempt_no,
                    'status': submissions[submission_id].status,
                    'submitted_at': submissions[submission_id].submitted_at,
                }
                for question, version, submission_id in submission_ids
            ],
        }, status=status.HTTP_201_CREATED)


class StudentSubmissionListView(APIView):
    """Daftar pengumpulan mahasiswa (status state machine P3)."""

    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request):
        submissions = list(
            Submission.objects.filter(student=request.user).order_by('-submitted_at')
        )

        version_ids = {s.question_version_id for s in submissions}
        versions = {
            v.id: v
            for v in QuestionVersion.objects.filter(id__in=version_ids)
        }
        question_ids = {v.question_id for v in versions.values()}
        questions = {
            q.id: q
            for q in Question.objects.filter(id__in=question_ids)
        }
        set_ids = {q.question_set_id for q in questions.values()}
        sets_map = {
            s.id: s
            for s in QuestionSet.objects.filter(id__in=set_ids).select_related('subject')
        }

        results = []
        for s in submissions:
            v = versions.get(s.question_version_id)
            q = questions.get(v.question_id) if v else None
            qs = sets_map.get(q.question_set_id) if q else None
            results.append({
                'id': str(s.id),
                'set_id': str(qs.id) if qs else None,
                'question_version_id': str(s.question_version_id),
                'question_prompt': v.prompt if v else None,
                'set_title': qs.title if qs else None,
                'set_code': qs.code if qs else None,
                'subject_id': str(qs.subject_id) if qs else None,
                'subject_name': qs.subject.name if qs else None,
                'attempt_no': s.attempt_no,
                'status': s.status,
                'answer_text': s.answer_text,
                'submitted_at': s.submitted_at,
            })

        return Response(results)


# ============================================================================
# SPRINT 4: STUDENT SUBMISSION GROUPING BY QUESTION SET
# ============================================================================


def _build_submission_set_groups(user):
    """Group a student's submissions per question set with comprehensive pedagogical diagnostics.

    Exposes rubric points breakdown, attempt-bound evaluations, model answer benchmarks,
    and verified misconceptions without leaking internal unvalidated AI hypotheses.
    """
    submissions = list(
        Submission.objects.filter(student=user).order_by('submitted_at')
    )
    if not submissions:
        return []

    version_ids = {s.question_version_id for s in submissions}
    versions = {
        v.id: v
        for v in QuestionVersion.objects.filter(id__in=version_ids).select_related('question')
    }
    question_ids = {v.question_id for v in versions.values()}
    questions = {q.id: q for q in Question.objects.filter(id__in=question_ids)}
    set_ids = {q.question_set_id for q in questions.values()}
    sets_map = {
        s.id: s
        for s in QuestionSet.objects.filter(id__in=set_ids).select_related('subject', 'topic')
    }

    # Fetch concept indicators for all versions in scope
    indicators_by_version = {}
    for ind in ConceptIndicator.objects.filter(question_version_id__in=version_ids).order_by('order_index'):
        indicators_by_version.setdefault(ind.question_version_id, []).append(ind)

    # Fetch current analyses and validations for every attempt
    sub_ids = [s.id for s in submissions]
    analyses = {
        a.submission_id: a
        for a in LlmAnalysis.objects.filter(submission_id__in=sub_ids, is_current=True)
    }
    validations = {
        v.analysis_id: v
        for v in Validation.objects.filter(
            analysis_id__in=[a.id for a in analyses.values()]
        ).select_related('lecturer')
    }

    published_rows = (
        QuestionVersion.objects.filter(
            is_published=True, question__question_set_id__in=set_ids
        )
        .values_list('question_id', 'question__question_set_id')
        .distinct()
    )
    published_count = {}
    for _qid, _sid in published_rows:
        published_count[_sid] = published_count.get(_sid, 0) + 1

    grouped = {}
    for s in submissions:
        v = versions.get(s.question_version_id)
        q = questions.get(v.question_id) if v else None
        if not q:
            continue
        grouped.setdefault(q.question_set_id, {}).setdefault(q.id, []).append(s)

    results = []
    for set_id, question_map in grouped.items():
        q_set = sets_map.get(set_id)
        if not q_set:
            continue

        status_counts = {}
        max_attempt = 0
        last_submitted = None
        questions_payload = []

        for qid, subs in question_map.items():
            q = questions[qid]
            attempts = []
            for s in sorted(subs, key=lambda x: x.attempt_no):
                status_counts[s.status] = status_counts.get(s.status, 0) + 1
                max_attempt = max(max_attempt, s.attempt_no)
                if last_submitted is None or s.submitted_at > last_submitted:
                    last_submitted = s.submitted_at

                analysis = analyses.get(s.id)
                validation = validations.get(analysis.id) if analysis else None
                evaluation_payload = None

                if s.status == 'VALIDATED' and validation:
                    v_indicators = indicators_by_version.get(s.question_version_id, [])
                    breakdown = analysis.concept_breakdown_json or {}
                    indicator_results = {
                        item.get('order_index'): item
                        for item in breakdown.get('indicators', [])
                    }

                    # Construct explainable rubric scoring breakdown
                    rubric_breakdown = []
                    for ind in v_indicators:
                        res = indicator_results.get(ind.order_index, {})
                        score_enum = res.get('score', 'MISSING')
                        credit_factor = 1.0 if score_enum == 'PRESENT' else (0.5 if score_enum == 'PARTIAL' else 0.0)
                        weight_float = float(ind.weight)
                        earned_points = round(weight_float * credit_factor * 100, 2)

                        rubric_breakdown.append({
                            'order_index': ind.order_index,
                            'label': ind.label,
                            'description': ind.description or '',
                            'max_weight_percent': round(weight_float * 100, 1),
                            'earned_points_percent': earned_points,
                            'status': score_enum,
                            'evidence': res.get('evidence', ''),
                        })

                    # Filter lecturer-confirmed misconceptions
                    confirmed_misconceptions = []
                    for m in breakdown.get('misconception_matches', []):
                        if m.get('lecturer_confirmed') is True:
                            confirmed_misconceptions.append({
                                'label': m.get('label'),
                                'reasoning': m.get('reasoning'),
                            })

                    final_score = float(validation.final_percentage) if validation.final_percentage is not None else float(analysis.percentage_correct)

                    evaluation_payload = {
                        'percentage_correct': final_score,
                        'is_score_modified': validation.is_score_modified,
                        'original_ai_score': float(validation.original_percentage),
                        'tier_level': validation.final_tier_level_snapshot or analysis.tier_level_snapshot,
                        'tier_label': validation.final_tier_label_snapshot or analysis.tier_label_snapshot,
                        'clinical_feedback': validation.final_feedback or analysis.explanation,
                        'confirmed_misconceptions': confirmed_misconceptions,
                        'rubric_breakdown': rubric_breakdown,
                        'suggested_materials': analysis.suggested_materials_json or [],
                        'validator_name': validation.lecturer.full_name,
                        'validated_at': validation.validated_at.isoformat() if validation.validated_at else None,
                    }

                attempts.append({
                    'submission_id': str(s.id),
                    'attempt_no': s.attempt_no,
                    'status': s.status,
                    'answer_text': s.answer_text,
                    'submitted_at': s.submitted_at,
                    'evaluation': evaluation_payload,
                })

            # Retrieve prompt and reference model answer from latest version
            latest_version = versions[sorted(subs, key=lambda x: x.attempt_no)[-1].question_version_id]
            questions_payload.append({
                'question_id': str(qid),
                'order_index': q.order_index,
                'version_id': str(latest_version.id),
                'version_number': latest_version.version_number,
                'prompt': latest_version.prompt,
                'model_answer': latest_version.model_answer,
                'answered': True,
                'attempts': attempts,
            })

        unanswered = Question.objects.filter(question_set_id=set_id).exclude(
            id__in=list(question_map.keys())
        ).order_by('order_index')
        for q in unanswered:
            v = QuestionVersion.objects.filter(
                question=q, is_published=True
            ).order_by('-version_number').first()
            if not v:
                continue
            questions_payload.append({
                'question_id': str(q.id),
                'order_index': q.order_index,
                'version_id': str(v.id),
                'version_number': v.version_number,
                'prompt': v.prompt,
                'model_answer': v.model_answer,
                'answered': False,
                'attempts': [],
            })

        questions_payload.sort(key=lambda item: item['order_index'])
        status_summary = next(iter(status_counts)) if len(status_counts) == 1 else 'MIXED'

        results.append({
            'set_id': str(q_set.id),
            'code': q_set.code,
            'title': q_set.title,
            'subject_id': str(q_set.subject_id),
            'subject_name': q_set.subject.name,
            'topic_id': str(q_set.topic_id) if q_set.topic_id else None,
            'topic_name': q_set.topic.name if q_set.topic else None,
            'question_count': published_count.get(set_id, 0),
            'answered_count': sum(1 for item in questions_payload if item['answered']),
            'total_attempts': sum(status_counts.values()),
            'max_attempt_no': max_attempt,
            'status_summary': status_summary,
            'status_counts': status_counts,
            'last_submitted_at': last_submitted,
            'questions': questions_payload,
        })

    results.sort(key=lambda item: item['last_submitted_at'], reverse=True)
    return results

class StudentSubmissionSetListView(APIView):
    """Daftar pengumpulan mahasiswa dikelompokkan per bank soal (UC-01)."""

    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response(_build_submission_set_groups(request.user))


class StudentSubmissionSetDetailView(APIView):
    """Rincian satu bank soal: seluruh pertanyaan, jawaban, dan evaluasi."""

    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        for group in _build_submission_set_groups(request.user):
            if group['set_id'] == str(pk):
                return Response(group)
        return Response(
            {'detail': 'Pengumpulan untuk bank soal ini tidak ditemukan.'},
            status=status.HTTP_404_NOT_FOUND,
        )


class LecturerSubmissionsView(APIView):
    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user

        if user.is_superuser:
            scoped = Submission.objects.all()
        else:
            lecturer_subjects = list(
                UserSubjectRole.objects.filter(user=user, role=UserSubjectRole.Role.LECTURER)
                .values_list('subject_id', flat=True)
            )
            if not lecturer_subjects:
                return Response({'detail': 'Akses ditolak.'}, status=status.HTTP_403_FORBIDDEN)
            scoped = Submission.objects.filter(subject_id__in=lecturer_subjects)

        status_param = request.query_params.get('status')
        if status_param:
            scoped = scoped.filter(status=status_param)
        subject_slug = request.query_params.get('subject_slug')
        if subject_slug:
            scoped = scoped.filter(subject__slug=subject_slug)

        submissions = list(scoped.select_related('student', 'subject').order_by('-submitted_at')[:300])

        analyses = {
            a.submission_id: a
            for a in LlmAnalysis.objects.filter(
                submission_id__in=[s.id for s in submissions], is_current=True
            )
        }
        validations = {
            v.analysis_id: v
            for v in Validation.objects.filter(
                analysis_id__in=[a.id for a in analyses.values()]
            )
        }
        versions = {
            v.id: v
            for v in QuestionVersion.objects.filter(
                id__in={s.question_version_id for s in submissions}
            ).select_related('question__question_set')
        }

        rows = []
        for s in submissions:
            version = versions.get(s.question_version_id)
            question = version.question if version else None
            qset = question.question_set if question else None
            analysis = analyses.get(s.id)
            validation = validations.get(analysis.id) if analysis else None
            rows.append({
                'id': str(s.id),
                'student_name': s.student.full_name,
                'student_email': s.student.email,
                'subject_slug': s.subject.slug,
                'subject_name': s.subject.name,
                'question_code': qset.code if qset else None,
                'question_title': qset.title if qset else None,
                'version_number': version.version_number if version else None,
                'prompt_preview': (version.prompt[:240] + '…') if version and version.prompt else '',
                'answer': s.answer_text,
                'status': s.status,
                'submitted_at': s.submitted_at.isoformat() if s.submitted_at else None,
                'score': float(analysis.percentage_correct) if analysis and analysis.percentage_correct is not None else None,
                'tier_label': analysis.tier_label_snapshot if analysis else None,
                'validation_status': validation.status if validation else None,
            })

        return Response({'submissions': rows})


class LecturerSubmissionDetailView(APIView):
    """Submission-first review detail; analysis is supplementary and optional."""

    authentication_classes = [TokenAuthentication]
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            submission = Submission.objects.select_related('student', 'subject').get(pk=pk)
        except Submission.DoesNotExist:
            return Response({'detail': 'Jawaban tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)

        if not is_lecturer_for_subject(request.user, submission.subject_id):
            return Response({'detail': 'Anda tidak berwenang mengakses jawaban ini.'}, status=status.HTTP_403_FORBIDDEN)

        try:
            version = QuestionVersion.objects.select_related('question__question_set').get(
                pk=submission.question_version_id
            )
        except QuestionVersion.DoesNotExist:
            return Response({'detail': 'Versi soal untuk jawaban ini tidak ditemukan.'}, status=status.HTTP_404_NOT_FOUND)

        analysis = LlmAnalysis.objects.filter(submission=submission, is_current=True).first()
        validation = Validation.objects.select_related('lecturer').filter(analysis=analysis).first() if analysis else None
        q_set = version.question.question_set

        return Response({
            'id': str(submission.id),
            'status': submission.status,
            'submitted_at': submission.submitted_at,
            'attempt_no': submission.attempt_no,
            'student': {
                'id': str(submission.student_id),
                'name': submission.student.full_name,
                'email': submission.student.email,
            },
            'subject': {'id': str(submission.subject_id), 'name': submission.subject.name},
            'set': {'id': str(q_set.id), 'code': q_set.code, 'title': q_set.title},
            'question': {
                'id': str(version.question_id),
                'version_id': str(version.id),
                'version_number': version.version_number,
                'prompt': version.prompt,
                'model_answer': version.model_answer,
                'reference_answers': [
                    {'id': str(answer.id), 'answer_key': answer.answer_key, 'text': answer.answer_text}
                    for answer in ReferenceAnswer.objects.filter(question_version=version).order_by('-is_primary', 'created_at')
                ],
                'indicators': [
                    {
                        'order_index': indicator.order_index,
                        'label': indicator.label,
                        'description': indicator.description or '',
                        'weight': str(indicator.weight),
                    }
                    for indicator in ConceptIndicator.objects.filter(question_version=version).order_by('order_index')
                ],
            },
            'answer_text': submission.answer_text,
            'current_analysis': {
                'id': str(analysis.id),
                'run_number': analysis.run_number,
                'percentage_correct': str(analysis.percentage_correct),
                'tier_level': analysis.tier_level_snapshot,
                'tier_label': analysis.tier_label_snapshot,
                'confidence': str(analysis.confidence),
                'created_at': analysis.created_at,
                'validation': {
                    'status': validation.status,
                    'lecturer_name': validation.lecturer.full_name,
                    'validated_at': validation.validated_at,
                } if validation else None,
            } if analysis else None,
        })
